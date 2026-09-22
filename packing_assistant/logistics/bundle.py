"""Bounded logistics handover bundles, with complete retained ledger history."""
from __future__ import annotations

from copy import deepcopy
import base64
import hashlib
import io
import json
import stat
import zipfile

from packing_assistant.runtime.cancel import check
from .records import MAX_RECORD, MAX_SOURCE, source_bytes, validate_record, digest

MAX_ZIP = 24 * 1024 * 1024


def _json(data):
    def pairs(items):
        out = {}
        for key, value in items:
            if key in out:
                raise ValueError("JSON 出现重复字段。")
            out[key] = value
        return out
    try:
        return json.loads(data, object_pairs_hook=pairs, parse_constant=lambda _: (_ for _ in ()).throw(ValueError("非有限数字")))
    except (UnicodeError, RecursionError, TypeError) as exc:
        raise ValueError("项目包 JSON 无效或嵌套过深。") from exc


def export_bundle(record):
    record = validate_record(record)
    raw = source_bytes(record)
    payload = deepcopy(record)
    payload["confirmation"] = None
    payload.pop("source_file")
    project = json.dumps(payload, ensure_ascii=False, allow_nan=False, separators=(",", ":")).encode()
    sha = hashlib.sha256(raw).hexdigest()
    files = {"project.json": project, f"sources/{sha}.bin": raw}
    manifest = {"schema": "civil.logistics.bundle.v1", "files": {name: {"sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)} for name, data in files.items()}}
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("manifest.json", json.dumps(manifest, ensure_ascii=False).encode())
        for name, data in files.items():
            check()
            archive.writestr(name, data)
    data = output.getvalue()
    if len(data) > MAX_ZIP:
        raise ValueError("项目包超过 24 MiB 上限。")
    return data


def import_bundle(data):
    if not isinstance(data, bytes) or not 0 < len(data) <= MAX_ZIP:
        raise ValueError("项目包为空或超过 24 MiB。")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            names = [i.filename for i in infos]
            if len(names) != 3 or len(set(names)) != 3 or "manifest.json" not in names or "project.json" not in names:
                raise ValueError("项目包必须且只能含清单、台账与一个原件。")
            if sum(i.file_size for i in infos) > MAX_RECORD + MAX_SOURCE + 65536:
                raise ValueError("项目包展开大小超限。")
            for item in infos:
                check()
                if item.flag_bits & 1 or item.compress_type not in {zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED} or stat.S_ISLNK(item.external_attr >> 16) or item.is_dir():
                    raise ValueError("项目包不允许链接、加密或未知压缩。")
                limit = 65536 if item.filename == "manifest.json" else MAX_RECORD if item.filename == "project.json" else MAX_SOURCE
                if not 0 < item.file_size <= limit:
                    raise ValueError("项目包成员大小无效。")
            manifest = _json(archive.read("manifest.json"))
            if not isinstance(manifest, dict) or set(manifest) != {"schema", "files"} or manifest["schema"] != "civil.logistics.bundle.v1" or not isinstance(manifest["files"], dict):
                raise ValueError("项目包清单无效。")
            rawfiles = {}
            if set(manifest["files"]) != set(names) - {"manifest.json"}:
                raise ValueError("清单与 ZIP 成员不一致。")
            for name, meta in manifest["files"].items():
                if not isinstance(meta, dict) or set(meta) != {"sha256", "bytes"}:
                    raise ValueError("清单字段无效。")
                content = archive.read(name)
                if type(meta["bytes"]) is not int or len(content) != meta["bytes"] or hashlib.sha256(content).hexdigest() != meta["sha256"]:
                    raise ValueError("项目包成员摘要或大小不匹配。")
                rawfiles[name] = content
            record = _json(rawfiles["project.json"])
            source = record["document"]["source"]
            expected = f"sources/{source['sha256']}.bin"
            if set(rawfiles) != {"project.json", expected} or len(expected) != 76 or any(c not in "0123456789abcdef" for c in source["sha256"]):
                raise ValueError("原件成员名无效，不接受任意路径。")
            raw = rawfiles[expected]
            if "source_file" in record:
                raise ValueError("项目包含重复原件记录。")
            record["source_file"] = {"data": base64.b64encode(raw).decode(), "bytes": len(raw), "sha256": source["sha256"]}
            # Never restore assertions or authorization from an untrusted handover.
            record["confirmation"] = None
            return validate_record(record)
    except (zipfile.BadZipFile, RuntimeError, KeyError, TypeError, AttributeError, RecursionError, OverflowError) as exc:
        raise ValueError("项目包损坏或结构无效，未创建项目。") from exc


def export_ledger(project, format):
    from .ledger import validate_document
    doc = validate_document(project["document"])
    if format == "json":
        return json.dumps(doc, ensure_ascii=False, allow_nan=False, indent=2).encode(), "application/json", "logistics-ledger.json"
    if format != "xlsx":
        raise ValueError("台账导出只支持 json / xlsx。")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    fields = ["id", "container_id", "package_id", "package_type", "material_id", "name", "spec", "package_count", "quantity", "units_per_package", "unit", "length_mm", "width_mm", "height_mm", "dimension_scope", "net_kg", "gross_kg", "weight_scope"]
    wb = Workbook()
    ws = wb.active
    ws.title = "台账"
    ws.append(fields)
    for row in doc["rows"]:
        check()
        ws.append([row.get(f, "UNSPECIFIED") for f in fields])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"  # do not create Excel formulas from untrusted filenames/cells
    for cell in ws[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="16384B")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 20
    sources = wb.create_sheet("来源与修订")
    sources.append(["行号", "字段", "原文", "位置与修订记录"])
    for row in doc["rows"]:
        for field, evidence in row.get("evidence", {}).items():
            sources.append([row["id"], field, str(evidence.get("raw", ""))[:32767], json.dumps(evidence, ensure_ascii=False)[:32767]])
    meta = wb.create_sheet("记录")
    meta.append(["项目", project["name"]])
    meta.append(["修订", project["revision"]])
    meta.append(["原件", doc["source"]["filename"]])
    meta.append(["SHA-256", doc["source"]["sha256"]])
    meta.append(["说明", "此 XLSX 是台账交换副本；完整历史与原件请使用 ZIP 项目包。确认不替代工程签认。"])
    for sheet in (sources, meta):
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "logistics-ledger.xlsx"
