"""Local Excel/CSV/PDF/OCR tables -> an evidence ledger, never a packing estimate."""
from __future__ import annotations

import csv
import hashlib
import json
from io import BytesIO, StringIO
from pathlib import Path
import re
from zipfile import ZipFile, BadZipFile

from packing_assistant.runtime.cancel import check
from .ledger import SCHEMA, UNSPECIFIED, FIELDS, NUMERIC_FIELDS, COUNT_FIELDS, MAX_ROWS, audit_document, validate_document

MAX_BYTES = 20 * 1024 * 1024
MAX_PAGES = 40
MAX_PIXELS = 25_000_000
MAX_CELLS = 200000
MAX_COLUMNS = 128
_LENGTH = {"mm": 1, "毫米": 1, "cm": 10, "厘米": 10, "m": 1000, "米": 1000, "in": 25.4, "inch": 25.4, "inches": 25.4, "英寸": 25.4, "ft": 304.8, "英尺": 304.8}
_WEIGHT = {"kg": 1, "kgs": 1, "千克": 1, "公斤": 1, "g": .001, "克": .001, "t": 1000, "吨": 1000, "lb": .45359237, "lbs": .45359237, "磅": .45359237}
_QUANTITY_UNITS = {"pcs", "pc", "ea", "件", "个", "套", "卷", "米", "m", "根", "支", "片", "块", "张", "roll", "rolls", "set", "sets", "pair", "pairs"}
_PACKAGE_TYPES = {"箱", "木箱", "纸箱", "铁箱", "铁框", "铁架", "框", "扎", "捆", "袋", "卷", "包", "托", "托盘", "件", "个", "pcs", "pc", "ea", "packages", "cartons", "cases", "case", "pallet", "pallets", "bundles", "bundle", "bags", "bag", "crates", "crate"}


def _norm(value):
    return re.sub(r"[\s_.,:()（）\[\]/\\#-]+", "", str(value or "").lower())


_ALIASES = {
    "container_id": ("集装箱号", "柜号", "container no", "container number", "container id", "ctn no"),
    "package_type": ("包装类型", "包装形式", "package type", "packing type"),
    "package_id": ("箱号", "包装号", "package id", "package no", "case no", "case number", "box no", "carton no"),
    "material_id": ("物料编号", "材料编号", "物料编码", "材料编码", "料号", "件号", "item no", "item code", "part no", "material id", "sku"),
    "name": ("name", "品名", "货物名称", "材料名称", "名称", "货物名称及规格", "品名及规格", "description", "item description", "commodity"),
    "spec": ("spec", "规格", "型号", "specification", "model"),
    "package_count": ("箱数", "包装数", "包装数量", "包装总箱数", "number of packages", "no of packages", "packages", "cases", "cartons", "package count"),
    "quantity": ("quantity", "qty", "件数", "总件数", "数量", "total quantity", "total pcs", "pcs", "货物件数"),
    "units_per_package": ("每箱件数", "每箱数量", "pcs per carton", "pcs per box", "quantity per package", "units per package"),
    "unit": ("单位", "unit", "uom", "unit of measure"),
    "length_mm": ("length", "长", "长度", "箱长", "package length", "item length", "length mm"),
    "width_mm": ("width", "宽", "宽度", "箱宽", "package width", "item width", "width mm"),
    "height_mm": ("height", "高", "高度", "箱高", "package height", "item height", "height mm"),
    "net_kg": ("net", "net weight", "n.w.", "净重", "net kg"),
    "gross_kg": ("gross", "gross weight", "g.w.", "毛重", "gross kg"),
    "dimension_scope": ("dimension scope", "尺寸范围", "尺寸口径"),
    "weight_scope": ("weight scope", "重量范围", "重量口径"),
    "dimensions": ("dimensions", "dimension", "尺寸", "长宽高", "l x w x h", "l*w*h", "l×w×h", "package dimensions", "箱外尺寸"),
}
_LOOKUP = {_norm(a): f for f, aliases in _ALIASES.items() for a in (*aliases, f)}


def _unit_in(text, units):
    tokens = re.findall(r"[a-zA-Z]+|[\u4e00-\u9fff]+", str(text).lower().replace("_", " "))
    matches = [token for token in tokens if token in units]
    return matches[-1] if matches and len(set(matches)) == 1 else None


def _scope(text, dimension=False):
    n = _norm(text)
    if n in ("package", "箱", "每箱", "包装") or any(t in n for t in ("perpackage", "perbox", "percarton", "percase", "每箱", "箱外", "箱长", "箱宽", "箱高", "packagelength", "packagewidth", "packageheight", "packagedimensions")):
        return "package"
    if n in ("item", "单件", "每件") or any(t in n for t in ("peritem", "perpiece", "单件", "每件", "itemlength", "itemwidth", "itemheight")):
        return "item"
    if not dimension and (n in ("row", "整行", "行合计") or any(t in n for t in ("rowtotal", "整行", "行合计", "totalnet", "totalgross", "总净重", "总毛重"))):
        return "row"
    return UNSPECIFIED


def _header_field(text):
    n = _norm(text)
    field = _LOOKUP.get(n)
    if field is None:
        stripped = re.sub(r"(?<![a-z])(?:mm|cm|kg|kgs|lbs|lb|inch|inches|ft|pcs|ea|m|g|t)(?![a-z])|毫米|厘米|千克|公斤|吨", "", text.lower().replace("_", " "))
        field = _LOOKUP.get(_norm(stripped))
    if field is None:
        stripped = re.sub(r"(?:per\s*(?:package|box|carton|case|item|piece)|每箱|每件|单件|整行|总|total)", "", text, flags=re.I)
        stripped = re.sub(r"(?<![a-z])(?:mm|cm|kg|kgs|lbs|lb|inch|inches|ft|m|g|t)(?![a-z])|毫米|厘米|千克|公斤|吨", "", stripped.lower().replace("_", " "))
        field = _LOOKUP.get(_norm(stripped))
    return field


def _header_candidates(raw):
    text = str(raw or "").strip()
    # Join wrapped words within each language, but never prefer one language over
    # a contradictory translation such as "净重 / Gross weight".
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]+", text))
    latin = re.sub(r"[\u4e00-\u9fff]+", " ", text)
    return {field for part in (text, chinese, latin) if (field := _header_field(part))}


def _header(raw):
    text = str(raw or "").strip()
    candidates = _header_candidates(text)
    field = next(iter(candidates)) if len(candidates) == 1 else None
    return field, _unit_in(text, _LENGTH if field in ("length_mm", "width_mm", "height_mm", "dimensions") else _WEIGHT), _scope(text, field in ("length_mm", "width_mm", "height_mm", "dimensions"))


def _suffix(raw):
    match = re.fullmatch(r"[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?\s*([A-Za-z\u4e00-\u9fff]+)", str(raw).strip())
    return match[1].lower() if match else ""


def _number(raw, field, unit):
    text = str(raw).strip()
    if not text or text.upper() in ("UNSPECIFIED", "N/A", "NA", "TBD") or text in ("-", "—", "待定", "待填"):
        return UNSPECIFIED, ""
    if isinstance(raw, bool):
        return UNSPECIFIED, "布尔值不是数量或尺寸。"
    match = re.fullmatch(r"([+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*([A-Za-z\u4e00-\u9fff]*)", text)
    if not match:
        return UNSPECIFIED, "数值格式不明确，未拆分范围或猜测数字。"
    value, suffix = float(match[1].replace(",", "")), match[2].lower()
    if value < 0 or value > 1e12:
        return UNSPECIFIED, "数值超出允许范围。"
    if field in COUNT_FIELDS:
        allowed = _PACKAGE_TYPES if field == "package_count" else _QUANTITY_UNITS
        if value != int(value) or (suffix and suffix not in allowed):
            return UNSPECIFIED, "件数、箱数与每箱件数须为明确的非负整数。"
        return int(value), ""
    units = _LENGTH if field.endswith("_mm") else _WEIGHT
    if suffix and suffix not in units:
        return UNSPECIFIED, "数值单位不受支持。"
    if suffix and unit and suffix != unit and units.get(suffix) != units.get(unit):
        return UNSPECIFIED, "单元格单位与表头单位冲突，请核对。"
    use = suffix or unit
    if use not in units:
        return UNSPECIFIED, "未明确标注单位，不按数值大小推测。"
    result = value * units[use]
    if result > 1e12:
        return UNSPECIFIED, "单位转换结果过大。"
    return round(result, 8), ""


def _cell(value, **location):
    return {"raw": "" if value is None else str(value), "source": location}


def _xlsx(data):
    try:
        with ZipFile(BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > 4096 or sum(i.file_size for i in infos) > 60 * 1024 * 1024 or any(i.flag_bits & 1 for i in infos):
                raise ValueError("Excel 解压体积过大或已加密")
    except BadZipFile as exc:
        raise ValueError("Excel 文件损坏") from exc
    import openpyxl
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    tables, cells = [], 0
    try:
        if len(workbook.worksheets) > 40:
            raise ValueError("Excel 工作表超过 40 张")
        for sheet in workbook.worksheets:
            check()
            if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError("Excel 行列超出限制，请裁剪有效表格")
            rows = []
            for row_index, row in enumerate(sheet.iter_rows(), 1):
                check()
                cells += len(row)
                if cells > MAX_CELLS:
                    raise ValueError("Excel 单元格超过限制")
                parsed = []
                for column_index, c in enumerate(row, 1):
                    value = c.value
                    if isinstance(value, int) and re.fullmatch(r"0+", c.number_format or ""):
                        value = str(value).zfill(len(c.number_format))
                    cell = _cell(value, sheet=sheet.title, row=row_index, column=column_index)
                    if c.data_type == "f":
                        cell["reason"] = "公式未求值；请核对并提供明确原值，不使用可能过期的缓存。"
                    parsed.append(cell)
                rows.append(parsed)
            tables.append({"rows": rows, "source": {"sheet": sheet.title}})
    finally:
        workbook.close()
    return tables, [], "openpyxl-native"


def _csv(data):
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        try:
            text = data.decode("gb18030")
        except UnicodeDecodeError as exc:
            raise ValueError("CSV 文本编码无法识别") from exc
    if "\x00" in text:
        raise ValueError("CSV 含二进制内容")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows, cells = [], 0
    for index, row in enumerate(csv.reader(StringIO(text), dialect), 1):
        check()
        cells += len(row)
        if index > 10000 or len(row) > MAX_COLUMNS or cells > MAX_CELLS:
            raise ValueError("CSV 行列超出限制")
        rows.append([_cell(value, row=index, column=j) for j, value in enumerate(row, 1)])
    return [{"rows": rows, "source": {}}], [], "csv-native"


def _page_kind(page, found):
    """Use an explicit title above tables, never the word INVOICE in a cargo cell."""
    bottom = min([page.height * .4, *[table.bbox[1] for table in found]])
    text = page.crop((0, 0, page.width, max(bottom, 1))).extract_text() or ""
    kinds = set()
    for line in text.splitlines():
        compact = re.sub(r"\s+", "", line).lower()
        if compact in ("packinglist", "packinglist装箱单", "装箱单packinglist") or re.match(r"^装箱单(?:$|网址|website)", compact):
            kinds.add("packing_list")
        if compact in ("invoice", "commercialinvoice", "发票invoice", "invoice发票") or re.match(r"^发票(?:$|网址|website)", compact):
            kinds.add("invoice")
        if compact in ("contract", "合同contract", "contract合同") or re.match(r"^(?:合同|暂时进出口合同协议)(?:$|网址|website)", compact):
            kinds.add("contract")
    return next(iter(kinds)) if len(kinds) == 1 else "ambiguous" if kinds else "unknown"


def _pdf_cells(table, page_number, table_index):
    """Retain a merged cell once and record only geometrically proven row spans."""
    values, geometry = table.extract(), table.rows
    x_edges = {edge for region in geometry for cell in region.cells if cell is not None for edge in (cell[0], cell[2])}
    rows = []
    for ri, (cells, region) in enumerate(zip(values, geometry), 1):
        rows.append([_cell(value, page=page_number, row=ri, column=ci + 1,
                           bbox=list(region.cells[ci] or table.bbox), coordinate_system="pdf_points")
                     for ci, value in enumerate(cells)])
    for ri, region in enumerate(geometry):
        check()
        for ci, bbox in enumerate(region.cells):
            if bbox is None or not str(rows[ri][ci]["raw"]).strip():
                continue
            if any(bbox[0] + .1 < x < bbox[2] - .1 for x in x_edges):
                rows[ri][ci]["reason"] = "单元格横跨多个字段列，未猜测该值属于哪一列。"
                continue
            covered = [ri]
            # A missing cell by itself is not evidence of merging. Its row must
            # lie entirely inside the anchor cell's vertical extent and columns
            # must align; a horizontal colspan is deliberately not inferred.
            for other in range(ri + 1, len(geometry)):
                other_cells = geometry[other].cells
                if ci >= len(other_cells) or other_cells[ci] is not None:
                    break
                present = [cell for cell in other_cells if cell is not None]
                if not present:
                    break
                top = min(cell[1] for cell in present)
                bottom = min(cell[3] for cell in present if cell[3] > top)
                if top < bbox[1] - .1 or bottom > bbox[3] + .1:
                    break
                covered.append(other)
            if len(covered) < 2:
                continue
            merge = {"id": f"P{page_number}T{table_index}R{ri + 1}C{ci + 1}",
                     "rows": [i + 1 for i in covered], "anchor_row": ri + 1,
                     "source": dict(rows[ri][ci]["source"])}
            for other in covered:
                rows[other][ci]["merge"] = merge
                if other != ri:
                    rows[other][ci]["source"] = dict(merge["source"])
    return rows


def _pdf(data, ocr_backend):
    try:
        import pdfplumber
    except ImportError as exc:
        raise ValueError("数字 PDF 表格解析需要安装 pdfplumber") from exc
    tables, issues, scan_pages = [], [], []
    with pdfplumber.open(BytesIO(data)) as pdf:
        if len(pdf.pages) > MAX_PAGES:
            raise ValueError("PDF 超过 40 页限制")
        page_info = []
        for number, page in enumerate(pdf.pages, 1):
            check()
            if page.width * page.height * (150 / 72) ** 2 > MAX_PIXELS:
                raise ValueError("PDF 页面栅格化像素超过限制")
            text = (page.extract_text() or "").strip()
            found = page.find_tables()
            kind = _page_kind(page, found) if text else "unknown"
            if text and not found:
                found = page.find_tables(table_settings={"vertical_strategy": "text", "horizontal_strategy": "text", "min_words_vertical": 2})
            page_info.append((number, page, text, found, kind))
        has_packing_title = any(kind == "packing_list" for *_, kind in page_info)
        for number, page, text, found, kind in page_info:
            check()
            if kind in ("invoice", "contract"):
                issues.append({"code": "excluded_non_packing_page", "severity": "info", "page": number,
                               "document_type": kind, "message": f"第 {number} 页明确为{'发票' if kind == 'invoice' else '合同'}，已排除，不重复计入箱单。"})
                continue
            if kind == "ambiguous" or (has_packing_title and kind == "unknown"):
                issues.append({"code": "ambiguous_document_page", "severity": "warning", "page": number,
                               "message": f"第 {number} 页单据类型不明确，未并入已识别的装箱单；请核对是否为续页。"})
                continue
            if not text:
                scan_pages.append(number)
                continue
            if not found:
                issues.append({"code": "no_table_on_page", "severity": "warning", "message": f"第 {number} 页有文字但未识别表格；没有把全文数字猜成物料。"})
            for table_index, table in enumerate(found, 1):
                tables.append({"rows": _pdf_cells(table, number, table_index),
                               "source": {"page": number, "bbox": list(table.bbox), "coordinate_system": "pdf_points"},
                               "document_type": kind})
    if scan_pages:
        if ocr_backend == "none":
            issues.append({"code": "ocr_required", "severity": "error", "message": "扫描页需要本地 OCR：" + ",".join(map(str, scan_pages))})
        else:
            from .ocr import extract_tables
            output = extract_tables(data, ".pdf", pages=scan_pages)
            tables.extend(output["tables"])
            issues.extend(output.get("issues", []))
    tables.sort(key=lambda t: (t.get("source", {}).get("page", 0), t.get("source", {}).get("bbox", [0, 0])[1]))
    return tables, issues, "pdfplumber+PP-StructureV3" if scan_pages and ocr_backend != "none" else "pdfplumber"


def _summary_label(raw):
    parts = [part for part in re.split(r"[\s/]+", str(raw).strip()) if part]
    chinese = "".join(re.findall(r"[\u4e00-\u9fff]+", str(raw)))
    latin = _norm(re.sub(r"[\u4e00-\u9fff]+", "", str(raw)))
    known = ("合计", "总计", "小计", "本页合计", "total", "subtotal", "grandtotal")
    return bool(parts) and (chinese or latin) and all(not part or part in known for part in (chinese, latin))


def _table_rows(tables, issues):
    rows, totals = [], []
    prior_header = None
    for table in tables:
        check()
        matrix = table.get("rows", [])
        if len(matrix) > 10000 or sum(len(r) for r in matrix) > MAX_CELLS:
            raise ValueError("表格规模超过限制")
        header, section_ids, source_rows = None, [], {}
        table_start = len(rows)
        # A headerless continuation is used only when explicitly labelled by the extractor.
        if table.get("continuation") and prior_header:
            header = prior_header
        for matrix_index, cells in enumerate(matrix, 1):
            check()
            if len(cells) > MAX_COLUMNS:
                raise ValueError("表格列数超过限制")
            if not any(str(c.get("raw", "")).strip() for c in cells):
                continue
            mapped = [_header(c.get("raw", "")) for c in cells]
            fields = [h[0] for h in mapped if h[0]]
            conflicts = any(len(_header_candidates(c.get("raw", ""))) > 1 for c in cells)
            if (len(set(fields)) >= 2 or (fields and conflicts)) and any(f in fields for f in ("name", "package_id", "material_id")):
                header = [(mapping, cell.get("raw", "")) for mapping, cell in zip(mapped, cells)]
                for cell in cells:
                    if len(_header_candidates(cell.get("raw", ""))) > 1:
                        issues.append({"code": "conflicting_header", "severity": "error", "source": cell.get("source", {}),
                                       "message": "同一表头的中英文对应不同字段，已拒绝该列，需人工核对。"})
                prior_header = header
                continue
            if header is None:
                continue
            row = {"id": f"R{len(rows) + 1:05d}", **dict.fromkeys(FIELDS, UNSPECIFIED), "evidence": {}}
            scopes = {"dimension_scope": set(), "weight_scope": set()}
            targets = [h[0][0] for h in header if h[0][0]]
            for index, ((field, unit, scope), raw_header) in enumerate(header):
                if not field or index >= len(cells):
                    continue
                cell = cells[index]
                raw = str(cell.get("raw", ""))
                evidence = {"raw": raw, "source": cell.get("source", {}), "header": raw_header}
                if cell.get("merge"):
                    evidence["_merge"] = cell["merge"]
                if cell.get("ocr"):
                    evidence["ocr"] = True
                if "confidence" in cell:
                    evidence["confidence"] = cell["confidence"]
                if targets.count(field) > 1:
                    evidence["reason"] = "同一字段对应多个列，需人工选择，未自动取第一列。"
                    row["evidence"][field] = evidence
                    continue
                if field == "dimensions":
                    pieces = re.split(r"\s*[x×*]\s*", raw.strip())
                    for f, part in zip(("length_mm", "width_mm", "height_mm"), pieces if len(pieces) == 3 else ["", "", ""]):
                        value, reason = _number(part, f, unit)
                        row[f] = value
                        row["evidence"][f] = {**evidence, **({"reason": reason or "尺寸须明确长×宽×高及单位。"} if value == UNSPECIFIED else {})}
                    if scope != UNSPECIFIED:
                        scopes["dimension_scope"].add(scope)
                    continue
                if cell.get("reason"):
                    value, reason = UNSPECIFIED, str(cell["reason"])
                elif field in NUMERIC_FIELDS:
                    value, reason = _number(raw, field, unit)
                elif field in ("dimension_scope", "weight_scope"):
                    value, reason = _scope(raw, field == "dimension_scope"), ""
                else:
                    value, reason = raw.strip() or UNSPECIFIED, ""
                row[field] = value
                if reason:
                    evidence["reason"] = reason
                row["evidence"][field] = evidence
                if field in ("length_mm", "width_mm", "height_mm", "net_kg", "gross_kg") and scope != UNSPECIFIED:
                    scopes["dimension_scope" if field.endswith("_mm") else "weight_scope"].add(scope)
            # A unit written beside a quantity is evidence, not a pcs default.
            for field, target, allowed in (("quantity", "unit", _QUANTITY_UNITS), ("package_count", "package_type", _PACKAGE_TYPES)):
                ev = row["evidence"].get(field, {})
                suffix = _suffix(ev.get("raw", ""))
                if suffix in allowed and row[field] != UNSPECIFIED:
                    existing = row.get(target, UNSPECIFIED)
                    if existing != UNSPECIFIED and _norm(existing) != _norm(suffix):
                        row[field] = UNSPECIFIED
                        row[target] = UNSPECIFIED
                        ev["reason"] = "数值后缀与独立单位或包装类型列冲突，未选择其中一个。"
                        row["evidence"][target] = {**row["evidence"].get(target, {}), "raw": str(existing), "source": ev.get("source", {}), "reason": ev["reason"]}
                    else:
                        row[target] = suffix
                        row["evidence"][target] = {**{k: v for k, v in ev.items() if k != "_merge"}, "derived_from": "explicit_value_suffix"}
            for key, inferred in scopes.items():
                if row[key] != UNSPECIFIED:
                    inferred.add(row[key])
                if len(inferred) == 1:
                    row[key] = next(iter(inferred))
                    if key not in row["evidence"]:
                        related = [e for f, e in row["evidence"].items() if f in (("length_mm", "width_mm", "height_mm") if key == "dimension_scope" else ("net_kg", "gross_kg"))]
                        row["evidence"][key] = {"raw": " / ".join(e.get("header", "") for e in related), "source": related[0]["source"] if related else table.get("source", {}), "derived_from": "explicit_headers"}
                elif len(inferred) > 1:
                    row[key] = UNSPECIFIED
                    row["evidence"][key] = {"raw": " / ".join(sorted(inferred)), "source": table.get("source", {}), "reason": "同一行出现冲突的尺寸或重量范围。"}
            first_written = next((str(c.get("raw", "")) for c in cells if str(c.get("raw", "")).strip()), "")
            summary = first_written if _summary_label(first_written) else None
            if summary:
                totals.append({"label": summary, "values": {f: row[f] for f in NUMERIC_FIELDS if type(row[f]) in (int, float)}, "source": next((c.get("source", {}) for c in cells if str(c.get("raw", "")).strip() == summary), table.get("source", {})), "row_ids": list(section_ids)})
                section_ids = []
                continue
            if not any(row[f] != UNSPECIFIED for f in ("name", "package_id", "material_id", *NUMERIC_FIELDS)):
                continue
            rows.append(row)
            source_rows[matrix_index] = row["id"]
            section_ids.append(row["id"])
            if len(rows) > MAX_ROWS:
                raise ValueError("物料行超过 5000 行限制")
        for row in rows[table_start:]:
            for field, evidence in row["evidence"].items():
                merge = evidence.pop("_merge", None)
                if merge:
                    members = [source_rows[i] for i in merge["rows"] if i in source_rows]
                    anchor = source_rows.get(merge["anchor_row"])
                    if len(members) > 1 and anchor in members and len(members) == len(merge["rows"]):
                        evidence["group"] = {"id": merge["id"], "row_ids": members, "anchor_row_id": anchor, "source": merge["source"]}
                    elif len(members) != len(merge["rows"]):
                        evidence["reason"] = "合并单元格跨越了非物料行，未猜测其汇总范围。"
                        row[field] = UNSPECIFIED
        if header is None and matrix:
            issues.append({"code": "unrecognized_table", "severity": "warning", "message": "有表格未找到明确物料表头，未猜测字段或跨页对应关系。"})
    return rows, totals


def parse_document(data: bytes, filename: str, ocr_backend: str = "auto") -> dict:
    check()
    if not isinstance(data, bytes) or not data or len(data) > MAX_BYTES:
        raise ValueError("文件须为 1 字节至 20 MB")
    if not isinstance(filename, str) or not filename or len(filename) > 255 or any(c in filename for c in ("/", "\\", "\x00")):
        raise ValueError("文件名必须是不含目录的名称")
    if ocr_backend == "off":
        ocr_backend = "none"
    if ocr_backend not in ("auto", "none", "paddle", "paddleocr"):
        raise ValueError("仅支持 none/auto/paddleocr 本地 OCR")
    suffix = Path(filename).suffix.lower()
    if suffix == ".json":
        def pairs(items):
            value = {}
            for key, item in items:
                if key in value:
                    raise ValueError("JSON 存在重复字段")
                value[key] = item
            return value
        try:
            doc = validate_document(json.loads(data.decode("utf-8-sig"), object_pairs_hook=pairs))
        except (UnicodeError, json.JSONDecodeError) as exc:
            raise ValueError("物流 JSON 文件损坏") from exc
        original = doc["source"]
        doc["source"] = {"filename": filename, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)}
        doc["extraction"] = {**doc["extraction"], "engine": "civil-logistics-json", "original_source": original, "local_only": True}
        doc["report"] = audit_document(doc)
        check()
        return doc
    if suffix in (".xlsx", ".xlsm"):
        tables, issues, engine = _xlsx(data)
    elif suffix in (".csv", ".tsv"):
        tables, issues, engine = _csv(data)
    elif suffix == ".pdf":
        tables, issues, engine = _pdf(data, ocr_backend)
    elif suffix in (".png", ".jpg", ".jpeg"):
        if ocr_backend == "none":
            raise ValueError("图片需要启用本地 OCR")
        from .ocr import extract_tables
        result = extract_tables(data, suffix)
        tables, issues, engine = result["tables"], result.get("issues", []), "PP-StructureV3-local"
    else:
        raise ValueError("支持 xlsx/csv/tsv、数字或扫描 PDF、PNG/JPG 和 civil.logistics.v1 JSON；不支持旧版 xls")
    rows, totals = _table_rows(tables, issues)
    doc = {"schema": SCHEMA, "source": {"filename": filename, "sha256": hashlib.sha256(data).hexdigest(), "bytes": len(data)},
           "document_type": UNSPECIFIED, "extraction": {"engine": engine, "issues": issues, "table_count": len(tables), "local_only": True}, "rows": rows, "totals": totals}
    doc = validate_document(doc)
    doc["report"] = audit_document(doc)
    check()
    return doc
