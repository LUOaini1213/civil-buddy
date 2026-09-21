"""Bounded schedule interchange; no imported schedule is silently recalculated.

Native JSON/CSV/XLSX preserve Civil Buddy's day-based model. MSPDI is a
documented subset, with loss notices and the imported source dates kept apart
from a future calculation. Optional MPXJ converts files in a disposable process.
"""
from __future__ import annotations

import csv
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation
import hashlib
import importlib.util
import io
import json
import math
import os
from pathlib import Path
import re
import subprocess
import sys
import tempfile
import time
import zipfile
from xml.etree import ElementTree as ET

SCHEMA = "civil-buddy.planning-exchange.v1"
MAX_BYTES = 8 * 1024 * 1024
MAX_EXPANDED = 32 * 1024 * 1024
MAX_TASKS = 250
MAX_RESOURCES = 32
MAX_XML_ELEMENTS = 50000
MAX_CELL = 32000
NS = "http://schemas.microsoft.com/project"
TASK_ID_FIELD = "188743731"  # Microsoft PjCustomField.pjCustomTaskText1.
RESOURCE_ID_FIELD = "205520904"  # Microsoft PjCustomField.pjCustomResourceText1.
TASK_ALIAS = "CivilBuddyTaskId"
RESOURCE_ALIAS = "CivilBuddyResourceId"
LINK_TYPES = {"0": "FF", "1": "FS", "2": "SF", "3": "SS"}
CSV_FIELDS = ["record_type", "id", "name", "duration", "progress", "parent_id", "dependencies", "resources",
              "actual_start", "actual_finish", "start", "end", "capacity", "metadata"]
TASK_FIELDS = CSV_FIELDS[1:12]


def _check():
    from packing_assistant.runtime.cancel import check
    check()


def _json(value):
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def _loads(value):
    def bad(value):
        raise ValueError("JSON 不允许 NaN 或 Infinity。")
    try:
        return json.loads(value, parse_constant=bad)
    except (ValueError, UnicodeError, RecursionError) as exc:
        raise ValueError("计划 JSON 无效或嵌套过深。") from exc


def _date(value, label):
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ValueError(f"{label} 必须是明确的 YYYY-MM-DD 日期。")
    try:
        parsed = date.fromisoformat(value)
    except ValueError as exc:
        raise ValueError(f"{label} 日期无效。") from exc
    if not 1900 <= parsed.year <= 2100:
        raise ValueError(f"{label} 超出支持日期范围。")
    return value


def _int(value, label):
    if isinstance(value, bool) or not re.fullmatch(r"-?\d+", str(value)):
        raise ValueError(f"{label} 必须是整数。")
    return int(value)


def _bounded(data):
    if not isinstance(data, bytes) or not data or len(data) > MAX_BYTES:
        raise ValueError("计划文件必须非空，且不超过 8 MiB。")
    _check()


def _notice(report, code, message, entity_id=None):
    row = {"code": code, "message": message, "severity": "warning"}
    if entity_id is not None:
        row["entity_id"] = str(entity_id)
    if row not in report:
        if len(report) >= 500:
            raise ValueError("计划未支持的语义超过 500 项，请先简化源文件。")
        report.append(row)


def _validated(plan):
    from .planning import validate_plan
    return validate_plan(plan)


def _jvm_path():
    """Prefer a configured or project-local JRE without changing process globals."""
    configured = os.environ.get("CIVIL_JAVA_HOME")
    if configured:
        for relative in ("bin/server/jvm.dll", "lib/server/libjvm.so", "lib/server/libjvm.dylib"):
            candidate = Path(configured) / relative
            if candidate.is_file():
                return str(candidate)
        raise OSError("CIVIL_JAVA_HOME 未指向有效 Java 运行时。")
    root = Path(__file__).resolve().parents[2]
    roots = [root]
    marker = root / ".git"
    if marker.is_file():
        content = marker.read_text(encoding="utf-8").strip()
        if content.startswith("gitdir: "):
            gitdir = (root / content[8:]).resolve()
            if gitdir.parent.name == "worktrees" and gitdir.parent.parent.name == ".git":
                roots.append(gitdir.parent.parent.parent)
    for checkout in roots:
        for candidate in sorted((checkout / ".tools/java/temurin21").glob("*/bin/server/jvm.dll")):
            if candidate.is_file():
                return str(candidate)
    import jpype
    return jpype.getDefaultJVMPath()


def capabilities():
    packages = all(importlib.util.find_spec(name) is not None for name in ("mpxj", "jpype"))
    jvm, reason = False, "未安装可选 mpxj / JPype1 依赖；尚未启用 MPP/P6 导入。"
    if packages:
        try:
            jvm = Path(_jvm_path()).is_file()
            reason = "" if jvm else "没有找到可用 JVM；不会自动安装系统 Java。"
        except (OSError, RuntimeError) as exc:
            reason = "没有找到可用 JVM；不会自动安装系统 Java。"
    return {"imports": ["json", "csv", "xlsx", "xml"], "exports": ["json", "csv", "xlsx", "xml"],
            "mpxj": {"available": packages and jvm, "packages_available": packages, "jvm_available": jvm,
                     "reason": reason, "formats": ["mpp", "xer", "pmxml"], "native_mpp_export": False}}


def _source_dates(value, ids):
    if value is None:
        return {}
    if not isinstance(value, dict) or set(value) - ids:
        raise ValueError("原计划日期引用了未知任务。")
    result = {}
    for key, dates in value.items():
        if not isinstance(dates, dict) or set(dates) - {"start", "end"}:
            raise ValueError("原计划日期结构无效。")
        result[key] = {field: _date(item, f"{key}.{field}") for field, item in dates.items() if item is not None}
        if result[key].get("start") and result[key].get("end") and result[key]["end"] < result[key]["start"]:
            raise ValueError(f"任务 {key} 原计划结束早于开始。")
    return result


def import_plan(data: bytes, filename: str) -> dict:
    """Return a validated draft, original dates and notices; never calls calculate."""
    _bounded(data)
    filename = Path(str(filename).replace("\\", "/")).name
    if len(filename) > 200:
        raise ValueError("计划文件名过长。")
    suffix = Path(filename).suffix.lower().lstrip(".")
    report = []
    if suffix == "json":
        raw = _loads(data)
        if not isinstance(raw, dict):
            raise ValueError("计划 JSON 必须是对象。")
        if "schema" in raw:
            if raw["schema"] != SCHEMA or set(raw) - {"schema", "plan", "original_dates", "report"}:
                raise ValueError("不支持此计划交换格式或字段。")
            plan, original = raw.get("plan"), raw.get("original_dates", {})
            if raw.get("report"):
                _notice(report, "prior_report", "源交换包携带历史说明；请同时保留原文件核对。")
        else:
            plan, original = raw, {}
    elif suffix == "csv":
        plan, original = _import_csv(data)
    elif suffix == "xlsx":
        plan, original = _import_xlsx(data)
    elif suffix in {"xml", "pmxml"}:
        root = _xml(data)
        if root.tag == f"{{{NS}}}Project":
            plan, original = _import_mspdi(root, report)
        else:
            converted = _mpxj(data, suffix)
            plan, original = _import_mspdi(_xml(converted), report)
            _notice(report, "mpxj_conversion", "P6 XML 经 MPXJ 转为 MSPDI 后导入；原文件保持不变，未调用原排程引擎。")
    elif suffix in {"mpp", "xer"}:
        converted = _mpxj(data, suffix)
        plan, original = _import_mspdi(_xml(converted), report)
        _notice(report, "mpxj_conversion", "文件经 MPXJ 转为 MSPDI 后导入；原文件保持不变，未调用 Project/P6 排程引擎。")
    else:
        raise ValueError("支持 JSON、CSV、XLSX、MSPDI XML；MPP/P6 需要可选 MPXJ 与 JVM。")
    normalized = _validated(plan)
    original = _source_dates(original, {task["id"] for task in normalized["tasks"]})
    _notice(report, "source_dates", "导入未重新排程。原计划日期单独保留；后续计算使用本工作台日粒度规则，可能不同于源软件。")
    _check()
    return {"plan": normalized, "original_dates": original, "report": report, "requires_confirmation": True,
            "source": {"filename": filename, "format": suffix, "sha256": hashlib.sha256(data).hexdigest()}}


def export_plan(plan: dict, result: dict, format: str) -> dict:
    """Export a successful calculation, rejecting stale result/recipe pairs."""
    from .planning import calculate
    normalized = _validated(plan)
    if isinstance(result, dict) and result.get("kind") == "resource":
        _validate_resource_result(normalized, result)
    else:
        expected = calculate(normalized)["result"]
        if not isinstance(result, dict) or result.get("tasks") != expected["tasks"]:
            raise ValueError("计算结果与计划不一致，请重新计算后导出。")
    original = {row["id"]: {"start": _date(row["start"], "计算开始"), "end": _date(row["end"], "计算结束")}
                for row in result["tasks"]}
    report = []
    if result.get("kind") == "resource":
        _notice(report, "resource_dates", "导出日期来自已验证的资源方案。重新导入后须明确选择计算方式；Project 或本工作台重新计算可能改变这些日期。")
    format = str(format).lower()
    if format == "json":
        data = _json({"schema": SCHEMA, "plan": normalized, "original_dates": original, "report": report}).encode()
        media = "application/json"
    elif format == "csv":
        data = _export_csv(normalized, original)
        media = "text/csv; charset=utf-8"
    elif format == "xlsx":
        data = _export_xlsx(normalized, original)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    elif format in {"xml", "mspdi"}:
        data = _export_mspdi(normalized, result)
        format, media = "xml", "application/xml"
        _notice(report, "mspdi_day_model", "按每天 8 小时输出日粒度 MSPDI，保留四类依赖及整数工作日时距；不包含小时班次、成本或原生 MPP。")
    else:
        raise ValueError("只支持 JSON、CSV、XLSX 和 MSPDI XML 导出，不支持原生 MPP/P6 写出。")
    if len(data) > MAX_BYTES:
        raise ValueError("导出超过 8 MiB 上限，请缩小计划。")
    _check()
    return {"data": data, "filename": "civil-planning." + format, "media_type": media, "report": report}


def _validate_resource_result(plan, result):
    from .planning import _calendar
    dates = _calendar(plan)
    source = {row["id"]: row for row in plan["tasks"]}
    children = {}
    for task in plan["tasks"]:
        if task["parent_id"]:
            children.setdefault(task["parent_id"], []).append(task["id"])
    if not isinstance(result.get("tasks"), list) or len(result["tasks"]) != len(source):
        raise ValueError("资源方案任务数量与原计划不一致。")
    rows = {}
    for row in result["tasks"]:
        if not isinstance(row, dict) or row.get("id") not in source or row["id"] in rows:
            raise ValueError("资源方案任务编号无效或重复。")
        ident = row["id"]; task = source[ident]
        start, finish = row.get("start_offset"), row.get("finish_offset")
        if type(start) is not int or type(finish) is not int or not 0 <= start <= finish <= len(dates) or start >= len(dates):
            raise ValueError("资源方案工作日范围无效。")
        if row.get("duration") != finish - start or row.get("start") != dates[start]:
            raise ValueError("资源方案工期与工作日日期映射不一致。")
        if ident not in children and row.get("end") != dates[finish - 1 if finish > start else start]:
            raise ValueError("资源方案工期与工作日日期映射不一致。")
        if any(row.get(key) != task[key] for key in ("name", "parent_id", "dependencies", "resources")):
            raise ValueError("资源方案输入与原计划不一致。")
        rows[ident] = row
    leaves = {key: task for key, task in source.items() if key not in children}
    for ident, task in leaves.items():
        row = rows[ident]
        if row["duration"] != task["duration"]:
            raise ValueError("资源方案修改了原计划工期。")
        for dep in task["dependencies"]:
            left = rows[dep["task_id"]]["finish_offset" if dep["type"][0] == "F" else "start_offset"]
            right = row["finish_offset" if dep["type"][1] == "F" else "start_offset"]
            if right < left + dep["lag"]:
                raise ValueError("资源方案不满足任务依赖。")
    for ident, members in children.items():
        if rows[ident]["start_offset"] != min(rows[key]["start_offset"] for key in members) or rows[ident]["finish_offset"] != max(rows[key]["finish_offset"] for key in members):
            raise ValueError("资源方案 WBS 汇总范围不一致。")
        if rows[ident].get("end") != max(rows[key]["end"] for key in members):
            raise ValueError("资源方案 WBS 汇总日期不一致。")
    for resource in plan["resources"]:
        events = {}
        for ident, task in leaves.items():
            demand = task["resources"].get(resource["id"], 0)
            start, finish = rows[ident]["start_offset"], rows[ident]["finish_offset"]
            if demand and finish > start:
                events[start] = events.get(start, 0) + demand
                events[finish] = events.get(finish, 0) - demand
        load = 0
        for offset in sorted(events):
            load += events[offset]
            if load > resource["capacity"]:
                raise ValueError("资源方案超过明确资源容量。")
    _check()


def _metadata(plan):
    return {"schema": SCHEMA, "start_date": plan["start_date"], "calendar": plan["calendar"], "csv_literals": "apostrophe-v1"}


def _task_row(task, original):
    return {**task, "dependencies": _json(task["dependencies"]), "resources": _json(task["resources"]),
            "parent_id": task["parent_id"] or "", "actual_start": task["actual_start"] or "",
            "actual_finish": task["actual_finish"] or "", **original.get(task["id"], {})}


def _safe_csv(value):
    value = str(value)
    if value.startswith("'") or value.lstrip().startswith(("=", "+", "-", "@")) or value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _export_csv(plan, original):
    stream = io.StringIO(newline="")
    writer = csv.DictWriter(stream, fieldnames=CSV_FIELDS, lineterminator="\r\n", extrasaction="raise")
    writer.writeheader()
    writer.writerow({"record_type": "plan", "metadata": _json(_metadata(plan))})
    for resource in plan["resources"]:
        writer.writerow({"record_type": "resource", **{k: _safe_csv(v) for k, v in resource.items()}})
    for task in plan["tasks"]:
        writer.writerow({"record_type": "task", **{k: _safe_csv(v) for k, v in _task_row(task, original).items()}})
    return ("\ufeff" + stream.getvalue()).encode("utf-8")


def _import_csv(data):
    try:
        reader = csv.DictReader(io.StringIO(data.decode("utf-8-sig"), newline=""))
        if reader.fieldnames != CSV_FIELDS:
            raise ValueError("CSV 表头必须与工作台交换模板一致，不能重复或缺失。")
        rows = list(reader)
    except (UnicodeError, csv.Error) as exc:
        raise ValueError("CSV 必须是有效 UTF-8 文件。") from exc
    if not rows or set(rows[0]) != set(CSV_FIELDS) or len(rows) > MAX_TASKS + MAX_RESOURCES + 1:
        raise ValueError("CSV 表头或数量无效，请使用工作台导出的交换模板。")
    if rows[0]["record_type"] != "plan":
        raise ValueError("CSV 首行必须包含计划日期和日历元数据。")
    metadata = _loads(rows[0]["metadata"])
    tasks, resources = [], []
    for row in rows[1:]:
        if None in row or any(value is None or len(value) > MAX_CELL for value in row.values()):
            raise ValueError("CSV 单元格无效或过大。")
        row = {key: (value[1:] if value.startswith("'") else value) for key, value in row.items()}
        if row["record_type"] == "task":
            tasks.append({key: row[key] for key in TASK_FIELDS})
        elif row["record_type"] == "resource":
            resources.append({key: row[key] for key in ("id", "name", "capacity")})
        else:
            raise ValueError("CSV 存在未知记录类型。")
    return _table_plan(metadata, tasks, resources)


def _table_plan(metadata, tasks, resources):
    if not isinstance(metadata, dict) or metadata.get("schema") != SCHEMA:
        raise ValueError("交换表缺少受支持的计划元数据。")
    plan = {"start_date": metadata.get("start_date"), "calendar": metadata.get("calendar"), "tasks": [], "resources": []}
    original = {}
    for row in tasks:
        identifier = str(row.get("id") or "")
        original[identifier] = {key: str(row[key]) for key in ("start", "end") if row.get(key)}
        plan["tasks"].append({"id": identifier, "name": row.get("name"), "duration": _int(row.get("duration"), "工期"),
                              "progress": _int(row.get("progress"), "进度"), "parent_id": row.get("parent_id") or None,
                              "actual_start": row.get("actual_start") or None, "actual_finish": row.get("actual_finish") or None,
                              "dependencies": _loads(row.get("dependencies", "[]")), "resources": _loads(row.get("resources", "{}"))})
    for row in resources:
        plan["resources"].append({"id": str(row.get("id") or ""), "name": row.get("name"), "capacity": _int(row.get("capacity"), "资源容量")})
    return plan, original


def _export_xlsx(plan, original):
    from openpyxl import Workbook
    book = Workbook(); metadata = book.active; metadata.title = "Plan"
    metadata.append(["metadata", _json(_metadata(plan))])
    tasks = book.create_sheet("Tasks"); tasks.append(TASK_FIELDS)
    for task in plan["tasks"]:
        row = _task_row(task, original)
        tasks.append([row[field] for field in TASK_FIELDS])
    resources = book.create_sheet("Resources"); resources.append(["id", "name", "capacity"])
    for resource in plan["resources"]:
        resources.append([resource[k] for k in ("id", "name", "capacity")])
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Literal text, including =,+,-,@; no formula execution.
        sheet.freeze_panes = "A2"
    buffer = io.BytesIO(); book.save(buffer); book.close()
    return buffer.getvalue()


def _import_xlsx(data):
    from openpyxl import load_workbook
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 256 or sum(info.file_size for info in entries) > MAX_EXPANDED:
                raise ValueError("XLSX 解压尺寸或条目数超过上限。")
            if any(info.flag_bits & 1 or info.filename.startswith("xl/externalLinks/") or "vbaProject" in info.filename for info in entries):
                raise ValueError("不接受加密、宏或外部工作簿链接。")
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except (zipfile.BadZipFile, KeyError, OSError, ET.ParseError) as exc:
        raise ValueError("XLSX 文件损坏或格式无效。") from exc
    try:
        if set(book.sheetnames) != {"Plan", "Tasks", "Resources"}:
            raise ValueError("XLSX 必须包含 Plan、Tasks、Resources 三张交换表。")
        tables = {}
        for sheet in book:
            if (sheet.max_row or 0) > MAX_TASKS + MAX_RESOURCES + 2 or (sheet.max_column or 0) > 32:
                raise ValueError("XLSX 行列数超过上限。")
            table = []
            for cells in sheet.iter_rows():
                if len(table) > MAX_TASKS + MAX_RESOURCES + 1 or len(cells) > 32:
                    raise ValueError("XLSX 实际行列数超过上限。")
                for cell in cells:
                    if cell.data_type == "f" or isinstance(cell.value, str) and len(cell.value) > MAX_CELL:
                        raise ValueError("交换表不接受公式或过大的单元格。")
                table.append([cell.value for cell in cells])
            tables[sheet.title] = table
        meta = tables["Plan"]
        if len(meta) != 1 or len(meta[0]) != 2 or meta[0][0] != "metadata":
            raise ValueError("Plan 元数据表无效。")
        def records(name, fields):
            rows = tables[name]
            if not rows or rows[0] != fields:
                raise ValueError(f"{name} 表头无效。")
            return [dict(zip(fields, row)) for row in rows[1:]]
        return _table_plan(_loads(meta[0][1]), records("Tasks", TASK_FIELDS), records("Resources", ["id", "name", "capacity"]))
    finally:
        book.close()


def _xml(data):
    from defusedxml.ElementTree import fromstring
    from defusedxml.common import DefusedXmlException
    try:
        root = fromstring(data, forbid_dtd=True, forbid_entities=True, forbid_external=True)
    except (ET.ParseError, DefusedXmlException, ValueError) as exc:
        raise ValueError("XML 损坏，或包含不允许的 DTD/实体。") from exc
    count, stack = 0, [(root, 0)]
    while stack:
        item, depth = stack.pop(); count += 1
        if count > MAX_XML_ELEMENTS or depth > 32:
            raise ValueError("XML 节点或嵌套超过上限。")
        if len(item.text or "") > MAX_CELL or any(len(value) > MAX_CELL for value in item.attrib.values()):
            raise ValueError("XML 单个字段过大。")
        stack.extend((child, depth + 1) for child in item)
    return root


def _tag(parent, name, value=None):
    element = ET.SubElement(parent, f"{{{NS}}}{name}")
    if value is not None:
        element.text = str(value)
    return element


def _find(parent, path):
    return parent.find("/".join(f"{{{NS}}}{part}" for part in path.split("/")))


def _all(parent, path):
    return parent.findall("/".join(f"{{{NS}}}{part}" for part in path.split("/")))


def _text(parent, path, default=None):
    element = _find(parent, path)
    return element.text.strip() if element is not None and element.text else default


def _xml_date(parent, field, required=False):
    value = _text(parent, field)
    if not value:
        if required:
            raise ValueError(f"MSPDI 缺少明确的 {field}。")
        return None
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:T\d\d:\d\d:\d\d(?:\.\d+)?)?", value):
        raise ValueError(f"MSPDI 的 {field} 日期格式无效。")
    return _date(value[:10], field)


def _duration(value, minutes, label):
    match = re.fullmatch(r"PT(?:(\d+(?:\.\d+)?)H)?(?:(\d+(?:\.\d+)?)M)?(?:(\d+(?:\.\d+)?)S)?", value or "")
    if not match or not any(match.groups()):
        raise ValueError(f"{label} 缺少明确工期或使用不支持的工期格式。")
    hours, mins, seconds = (Decimal(v or "0") for v in match.groups())
    days = (hours * 60 + mins + seconds / 60) / minutes
    if days != days.to_integral_value():
        raise ValueError(f"{label} 含不足整工作日工期；本工作台不会舍入。")
    return int(days)


def _resource_units(value, label):
    """MSPDI 1.0 means a full resource (100%), not a 100-unit pool."""
    units = _decimal(value, label)
    if not 1 <= units <= 10000 or units != units.to_integral_value():
        raise ValueError(f"{label} 须为 1–10000 个整数完整资源；MSPDI 1.0 表示 1 个资源（100%），不支持部分资源，不会缩放或舍入。")
    return int(units)


def _decimal(value, label):
    try:
        if not isinstance(value, str) or len(value) > 100:
            raise ValueError()
        result = Decimal(value)
        if not result.is_finite():
            raise ValueError()
        return result
    except (ValueError, InvalidOperation) as exc:
        raise ValueError(f"{label} 必须是有限数值。") from exc


def _calendar(root, report):
    uid = _text(root, "CalendarUID")
    calendars = _all(root, "Calendars/Calendar")
    matches = [item for item in calendars if _text(item, "UID") == uid]
    if not uid or len(matches) != 1:
        raise ValueError("MSPDI 缺少明确的项目日历。")
    calendar = matches[0]
    if _text(calendar, "BaseCalendarUID", "-1") not in {"-1", "0"}:
        raise ValueError("暂不支持继承日历，请先在源软件展开基础日历。")
    weekdays, holidays, seen, legacy_exceptions = [], [], set(), []
    for item in _all(calendar, "WeekDays/WeekDay"):
        number = _int(_text(item, "DayType"), "星期类型")
        if number == 0:
            # MSPDI 2003 and MPXJ use DayType=0 for dated exceptions.
            legacy_exceptions.append(item)
            continue
        if not 1 <= number <= 7 or number in seen:
            raise ValueError("MSPDI 必须明确列出互不重复的 7 个星期日类型。")
        seen.add(number)
        if _text(item, "DayWorking") == "1":
            weekdays.append((number + 5) % 7)
        elif _text(item, "DayWorking") != "0":
            raise ValueError("MSPDI 日历缺少工作日状态。")
    if len(seen) != 7:
        raise ValueError("MSPDI 日历未完整列明一周，请先补齐日历。")
    for item in [*_all(calendar, "Exceptions/Exception"), *legacy_exceptions]:
        if _text(item, "Type", "1") != "1" or _text(item, "DayWorking") != "0":
            raise ValueError("暂不支持循环节假日或额外工作日，请先展开成非工作日期。")
        period = _find(item, "TimePeriod")
        if period is None:
            raise ValueError("日历例外缺少日期区间。")
        start = date.fromisoformat(_xml_date(period, "FromDate", True))
        end = date.fromisoformat(_xml_date(period, "ToDate", True))
        span = (end - start).days
        if not 0 <= span <= 3660 or len(holidays) + span + 1 > 3660:
            raise ValueError("节假日区间无效或过大。")
        holidays.extend((start + timedelta(days=offset)).isoformat() for offset in range(span + 1))
    _notice(report, "day_granularity", "MSPDI 工作时间按项目 MinutesPerDay 换算为整数工作日；小时班次和日内起止时间不保留。")
    return {"weekdays": sorted(weekdays), "holidays": sorted(set(holidays))}, uid


def _import_mspdi(root, report):
    if root.tag != f"{{{NS}}}Project":
        raise ValueError("不是 Microsoft Project MSPDI XML。")
    start = _xml_date(root, "StartDate", True)
    minutes = _int(_text(root, "MinutesPerDay"), "每天工作分钟")
    if not 1 <= minutes <= 1440:
        raise ValueError("MinutesPerDay 超出范围。")
    calendar, calendar_uid = _calendar(root, report)
    aliases = {_text(item, "FieldID"): _text(item, "Alias") for item in _all(root, "ExtendedAttributes/ExtendedAttribute")}
    def identifier(item, prefix, field, alias):
        uid = _text(item, "UID")
        if uid is None or not re.fullmatch(r"\d+", uid):
            raise ValueError("MSPDI 对象缺少有效 UID。")
        if aliases.get(field) == alias:
            found = [_text(row, "Value") for row in _all(item, "ExtendedAttribute") if _text(row, "FieldID") == field]
            if len(found) == 1 and found[0]:
                return uid, found[0]
        return uid, prefix + uid
    source_tasks = [item for item in _all(root, "Tasks/Task") if _text(item, "UID") != "0" and _text(item, "IsNull", "0") != "1"]
    source_resources = [item for item in _all(root, "Resources/Resource") if _text(item, "UID") != "0" and _text(item, "IsNull", "0") != "1"]
    if len(source_tasks) > MAX_TASKS or len(source_resources) > MAX_RESOURCES:
        raise ValueError("MSPDI 任务或资源超过上限。")
    task_ids, resource_ids, tasks, resources, original, parents = {}, {}, [], [], {}, []
    for item in source_resources:
        uid, key = identifier(item, "R", RESOURCE_ID_FIELD, RESOURCE_ALIAS)
        if uid in resource_ids:
            raise ValueError("MSPDI 资源 UID 重复。")
        if _text(item, "Type", "1") != "1":
            raise ValueError(f"资源 {key} 不是工作资源；当前不支持材料或成本资源，不能将其按工作资源容量排程。")
        if _text(item, "CalendarUID", "-1") not in {"-1", calendar_uid}:
            raise ValueError(f"资源 {key} 使用独立资源日历；当前只支持明确的统一项目日历，不能忽略后继续排程。")
        resource_ids[uid] = key
        resources.append({"id": key, "name": _text(item, "Name"), "capacity": _resource_units(_text(item, "MaxUnits"), f"资源 {key} 容量")})
    pending = []
    for item in source_tasks:
        _check()
        uid, key = identifier(item, "T", TASK_ID_FIELD, TASK_ALIAS)
        if uid in task_ids:
            raise ValueError("MSPDI 任务 UID 重复。")
        task_ids[uid] = key
        level = _int(_text(item, "OutlineLevel"), "WBS 层级")
        if not 1 <= level <= 32 or level > len(parents) + 1:
            raise ValueError(f"任务 {key} 的 WBS 层级跳跃或无效。")
        parents = parents[:level - 1]
        parent_id = parents[-1] if parents else None
        summary = _text(item, "Summary", "0") == "1"
        duration = 0 if summary else _duration(_text(item, "Duration"), minutes, f"任务 {key}")
        if _text(item, "Milestone", "0") == "1" and duration != 0:
            raise ValueError(f"里程碑 {key} 的工期不为零。")
        task = {"id": key, "name": _text(item, "Name"), "duration": duration,
                "progress": _int(_text(item, "PercentComplete", "0"), "进度"), "parent_id": parent_id,
                "actual_start": _xml_date(item, "ActualStart"), "actual_finish": _xml_date(item, "ActualFinish"),
                "dependencies": [], "resources": {}}
        if summary and (task["progress"] or task["actual_start"] or task["actual_finish"]):
            _notice(report, "summary_derived", "源汇总进度/实际日期不作为独立输入，改由子任务汇总；原文件保持不变。", key)
            task.update(progress=0, actual_start=None, actual_finish=None)
        tasks.append(task); parents.append(key)
        original[key] = {field: value for field, value in (("start", _xml_date(item, "Start")), ("end", _xml_date(item, "Finish"))) if value}
        if parent_id and not pending[-1][1] and level > pending[-1][2]:
            raise ValueError("WBS 子任务的父任务未标记为汇总任务。")
        pending.append((item, summary, level, task))
        for field, normal in (("ConstraintType", "0"), ("Manual", "0"), ("EffortDriven", "0"), ("CalendarUID", "-1")):
            value = _text(item, field, normal)
            if value != normal and not (field == "CalendarUID" and value == calendar_uid):
                _notice(report, "unsupported_" + field, f"源字段 {field}={value} 不参与本工作台重排；原日期已保留。", key)
        for field in ("Baseline", "TimephasedData", "Cost", "Deadline", "SplitParts"):
            if _find(item, field) is not None:
                _notice(report, "unsupported_" + field, f"源字段 {field} 未映射到简化计划。", key)
        duration_format = _int(_text(item, "DurationFormat", "7"), "工期单位") % 32
        if duration_format in {4, 6, 8, 10, 12, 20}:
            raise ValueError(f"任务 {key} 使用连续日历时间工期，不能当成工作日工期。")
    for item, summary, level, task in pending:
        for link in _all(item, "PredecessorLink"):
            predecessor = _text(link, "PredecessorUID")
            kind = _text(link, "Type", "1")
            if predecessor not in task_ids or kind not in LINK_TYPES:
                raise ValueError(f"任务 {task['id']} 依赖未知任务或关系类型。")
            lag = _decimal(_text(link, "LinkLag", "0"), "依赖时距") / (minutes * 10)
            if not lag.is_finite() or lag != lag.to_integral_value():
                raise ValueError(f"任务 {task['id']} 的时距不是整工作日。")
            lag_format = _int(_text(link, "LagFormat", "7"), "时距单位") % 32
            if lag_format in {4, 6, 8, 10, 12, 20} and lag:
                raise ValueError(f"任务 {task['id']} 使用连续日历时距，不能当成工作日时距。")
            task["dependencies"].append({"task_id": task_ids[predecessor], "type": LINK_TYPES[kind], "lag": int(lag)})
    lookup = {task["id"]: task for task in tasks}
    for item in _all(root, "Assignments/Assignment"):
        task_uid, resource_uid = _text(item, "TaskUID"), _text(item, "ResourceUID")
        if resource_uid in {"-65535", "-1"}:
            _notice(report, "unassigned_resource", "源文件未分配资源的占位记录未导入。", task_uid)
            continue
        if task_uid not in task_ids or resource_uid not in resource_ids:
            raise ValueError("资源分配引用未知任务或资源。")
        resources_map = lookup[task_ids[task_uid]]["resources"]
        key = resource_ids[resource_uid]
        if key in resources_map:
            raise ValueError("同一任务与资源有重复分配，不能无提示合并。")
        resources_map[key] = _resource_units(_text(item, "Units"), f"任务 {task_ids[task_uid]} 的资源 {key} 分配")
    return {"start_date": start, "calendar": calendar, "tasks": tasks, "resources": resources}, original


def _export_mspdi(plan, result):
    ET.register_namespace("", NS)
    root = ET.Element(f"{{{NS}}}Project")
    for key, value in (("SaveVersion", 14), ("Name", "Civil Buddy planning"), ("ScheduleFromStart", 1),
                       ("StartDate", plan["start_date"] + "T08:00:00"), ("MinutesPerDay", 480),
                       ("MinutesPerWeek", 480 * len(plan["calendar"]["weekdays"])), ("CalendarUID", 1)):
        _tag(root, key, value)
    attributes = _tag(root, "ExtendedAttributes")
    for field, alias in ((TASK_ID_FIELD, TASK_ALIAS), (RESOURCE_ID_FIELD, RESOURCE_ALIAS)):
        attribute = _tag(attributes, "ExtendedAttribute")
        _tag(attribute, "FieldID", field); _tag(attribute, "FieldName", "Text1"); _tag(attribute, "Alias", alias)
    cal = _tag(_tag(root, "Calendars"), "Calendar")
    for key, value in (("UID", 1), ("Name", "Civil Buddy calendar"), ("IsBaseCalendar", 1), ("BaseCalendarUID", -1)):
        _tag(cal, key, value)
    weekdays = _tag(cal, "WeekDays")
    for number in range(1, 8):
        day = _tag(weekdays, "WeekDay"); _tag(day, "DayType", number)
        working = (number + 5) % 7 in plan["calendar"]["weekdays"]
        _tag(day, "DayWorking", int(working))
        if working:
            times = _tag(day, "WorkingTimes")
            for start, end in (("08:00:00", "12:00:00"), ("13:00:00", "17:00:00")):
                time_range = _tag(times, "WorkingTime"); _tag(time_range, "FromTime", start); _tag(time_range, "ToTime", end)
    exceptions = _tag(cal, "Exceptions")
    for holiday in plan["calendar"]["holidays"]:
        exception = _tag(exceptions, "Exception"); _tag(exception, "EnteredByOccurrences", 0)
        period = _tag(exception, "TimePeriod"); _tag(period, "FromDate", holiday + "T00:00:00"); _tag(period, "ToDate", holiday + "T23:59:00")
        _tag(exception, "Name", holiday); _tag(exception, "Type", 1); _tag(exception, "DayWorking", 0)
    children = {}
    for task in plan["tasks"]:
        children.setdefault(task["parent_id"], []).append(task)
    ordered = []
    def visit(parent, level, prefix):
        for index, task in enumerate(children.get(parent, []), 1):
            outline = prefix + [index]; ordered.append((task, level, ".".join(map(str, outline))))
            visit(task["id"], level + 1, outline)
    visit(None, 1, [])
    ids = {task["id"]: index for index, (task, _, _) in enumerate(ordered, 1)}
    resource_ids = {row["id"]: index for index, row in enumerate(plan["resources"], 1)}
    calculated = {row["id"]: row for row in result["tasks"]}
    tasks = _tag(root, "Tasks")
    for task, level, outline in ordered:
        row = calculated[task["id"]]; item = _tag(tasks, "Task")
        for key, value in (("UID", ids[task["id"]]), ("ID", ids[task["id"]]), ("Name", task["name"]),
                           ("OutlineNumber", outline), ("OutlineLevel", level), ("Summary", int(bool(children.get(task["id"])))),
                           ("Milestone", int(not children.get(task["id"]) and task["duration"] == 0)),
                           ("Start", row["start"] + "T08:00:00"), ("Finish", row["end"] + ("T08:00:00" if task["duration"] == 0 and not children.get(task["id"]) else "T17:00:00")),
                           ("Duration", f"PT{row['duration'] * 8}H0M0S"), ("DurationFormat", 7),
                           ("PercentComplete", task["progress"]), ("CalendarUID", 1)):
            _tag(item, key, value)
        for field, key in (("ActualStart", "actual_start"), ("ActualFinish", "actual_finish")):
            if task[key]:
                _tag(item, field, task[key] + ("T08:00:00" if key == "actual_start" else "T17:00:00"))
        attribute = _tag(item, "ExtendedAttribute"); _tag(attribute, "FieldID", TASK_ID_FIELD); _tag(attribute, "Value", task["id"])
        for dependency in task["dependencies"]:
            link = _tag(item, "PredecessorLink"); _tag(link, "PredecessorUID", ids[dependency["task_id"]])
            _tag(link, "Type", next(key for key, value in LINK_TYPES.items() if value == dependency["type"]))
            _tag(link, "LinkLag", dependency["lag"] * 4800); _tag(link, "LagFormat", 7)
    resources = _tag(root, "Resources")
    for resource in plan["resources"]:
        item = _tag(resources, "Resource")
        for key, value in (("UID", resource_ids[resource["id"]]), ("ID", resource_ids[resource["id"]]),
                           ("Name", resource["name"]), ("Type", 1), ("MaxUnits", float(resource["capacity"]))):
            _tag(item, key, value)
        attribute = _tag(item, "ExtendedAttribute"); _tag(attribute, "FieldID", RESOURCE_ID_FIELD); _tag(attribute, "Value", resource["id"])
    assignments = _tag(root, "Assignments"); index = 0
    for task in plan["tasks"]:
        for resource, units in task["resources"].items():
            index += 1; item = _tag(assignments, "Assignment")
            for key, value in (("UID", index), ("TaskUID", ids[task["id"]]), ("ResourceUID", resource_ids[resource]), ("Units", float(units))):
                _tag(item, key, value)
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


def _mpxj(data, suffix, timeout=45):
    state = capabilities()["mpxj"]
    if not state["available"]:
        raise ImportError(state["reason"])
    if os.environ.get("CIVIL_OS_SANDBOX_POLICY"):
        raise PermissionError("系统级沙箱不允许启动 MPXJ 子进程；未放宽沙箱。")
    with tempfile.TemporaryDirectory(prefix="civil-planning-import-") as directory:
        source, target = Path(directory) / ("source." + suffix), Path(directory) / "converted.xml"
        source.write_bytes(data)
        env = {key: value for key, value in os.environ.items() if not key.endswith("API_KEY") and key != "CIVIL_TOKEN"}
        env.update(PYTHON_DOTENV_DISABLED="1", PYTHONUTF8="1")
        process = subprocess.Popen([sys.executable, "-m", "packing_assistant.engineering.planning_exchange", "_convert", str(source), str(target)],
                                   cwd=Path(__file__).resolve().parents[2], env=env, stdin=subprocess.DEVNULL,
                                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        started = time.monotonic()
        try:
            while process.poll() is None:
                _check()
                if time.monotonic() - started > timeout:
                    raise TimeoutError("MPXJ 导入超时；已停止本次转换。")
                time.sleep(.05)
            _check()
            if process.returncode or not target.is_file() or target.stat().st_size > MAX_BYTES:
                raise ValueError("MPXJ 未生成有效的受限 MSPDI 文件；请用源软件导出 XML 后重试。")
            return target.read_bytes()
        finally:
            if process.poll() is None:
                process.kill()
            process.wait()


def _convert(source, target):
    """Child-only fixed conversion, using installed official MPXJ APIs."""
    import jpype
    import jpype.imports  # noqa: F401
    import mpxj  # noqa: F401; registers MPXJ jars before starting the JVM.
    if Path(source).stat().st_size > MAX_BYTES:
        raise ValueError("Input too large")
    jpype.startJVM(_jvm_path(), "-Xmx384m", "-XX:ActiveProcessorCount=1")
    try:
        from org.mpxj.reader import UniversalProjectReader
        from org.mpxj.mspdi import MSPDIWriter
        projects = UniversalProjectReader().readAll(str(source))
        if projects is None or projects.size() != 1:
            raise ValueError("请从源软件导出单个项目；不能静默选择多项目包的第一个计划。")
        project = projects.get(0)
        if project.getTasks().size() > MAX_TASKS + 1 or project.getResources().size() > MAX_RESOURCES + 1:
            raise ValueError("Unsupported or oversized project")
        MSPDIWriter().write(project, str(target))
    finally:
        jpype.shutdownJVM()


if __name__ == "__main__":
    if len(sys.argv) != 4 or sys.argv[1] != "_convert":
        raise SystemExit(2)
    _convert(sys.argv[2], sys.argv[3])
