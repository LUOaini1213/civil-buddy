"""Bounded session attachments using the Rust workbench metadata contract.

Original bytes and extracted text live under demo/data/uploads/<session>.
Metadata is published last, so readers never see an incomplete attachment.
"""

from __future__ import annotations

import json
import re
import time
from io import BytesIO
from pathlib import Path
from threading import RLock
from typing import Iterable
from uuid import uuid4
from xml.etree import ElementTree
from zipfile import ZipFile

from packing_assistant.sandbox import assert_open, assert_write, guarded_write_bytes, guarded_write_text
from packing_assistant.document_text import csv_text, docx_document_text, table_markdown

UPLOAD_ROOT = Path(__file__).resolve().parent / "data" / "uploads"
MAX_BYTES = 20 * 1024 * 1024
MAX_REQUEST_BYTES = 25 * 1024 * 1024
MAX_TEXT_CHARS = 200_000
MAX_FILES = 12
INJECT_CHARS = 60_000
MAX_ARCHIVE_BYTES = 60 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 4096
ALLOWED_EXT = frozenset({"pdf", "docx", "xlsx", "txt", "md", "csv", "json", "log"})
#: Refused attachments that gave no text, one JSON object per line. Not .json/.txt/.bin: those three are
#: the attachment records themselves (strict_documents, the Rust workbench's listing).
UNREADABLE_LOG = "unreadable.jsonl"
MAX_UNREADABLE = 24
_SESSION_RE = re.compile(r"[A-Za-z0-9-][A-Za-z0-9_-]{3,31}\Z")
_ID_RE = re.compile(r"[a-f0-9]{12}\Z")
_LOCK = RLock()


class UploadError(ValueError):
    """An attachment validation/extraction failure suitable for a 400 response."""


class UploadTooLarge(UploadError):
    """An attachment/request limit failure suitable for a 413 response."""


class UploadUnreadable(UploadError):
    """A file of an accepted type that gave no usable text: damaged, encrypted, a scan with no text
    layer, empty. Refused like any other bad upload - and remembered, because a bid check run later
    must be able to say "投标文件.pdf was given and could not be read" instead of "no response given"."""

    def __init__(self, message: str, *, name: str = "", kind: str = "", size: int = 0) -> None:
        super().__init__(message)
        self.name, self.kind, self.size = name, kind, size


def safe_session_id(session: str) -> str:
    if not isinstance(session, str) or not _SESSION_RE.fullmatch(session):
        raise UploadError("session_id 无效：需 4–32 位 ASCII 字母、数字、连字符或下划线，不能以下划线开头")
    return session


def safe_filename(filename: str) -> str:
    # Browser filenames may contain either platform's separators. Never retain
    # directory components, and preserve the extension when truncating a name.
    raw = str(filename or "").replace("\\", "/").rsplit("/", 1)[-1]
    cleaned = re.sub(r"[^A-Za-z0-9.\-_（）\u4e00-\u9fff]", "_", raw)
    if not cleaned:
        return "upload.bin"
    suffix = Path(cleaned).suffix
    if len(cleaned) > 80:
        cleaned = cleaned[:80 - len(suffix)] + suffix
    return cleaned


def _directory(session: str) -> Path:
    sid = safe_session_id(session)
    root = UPLOAD_ROOT.resolve()
    candidate = root / sid
    resolved = candidate.resolve()
    if resolved != candidate:
        raise UploadError("附件会话目录不允许链接跳转")
    try:
        resolved.relative_to(root)
    except ValueError as exc:
        raise UploadError("附件路径越界") from exc
    return resolved


def _path(directory: Path, filename: str) -> Path:
    candidate = directory / filename
    if candidate.resolve() != candidate:
        raise UploadError("附件文件不允许链接跳转")
    return candidate


#: under a text that was longer than the limit: the same words office_job writes under a job file it cut
CUT_NOTE = "（未读完）只读了前 {n} 个字符，后面的内容未参与解析"


def _collapse(text: str) -> str:
    lines: list[str] = []
    used = 0
    cut = False
    noted = MAX_TEXT_CHARS >= 1024   # a limit too small to hold the note (tests) keeps the bare cut
    limit = MAX_TEXT_CHARS - 64 if noted else MAX_TEXT_CHARS   # room for the note: a cut attachment says so, in its own text
    for line in text.splitlines():
        line = line.strip()
        if not line and (not lines or not lines[-1]):
            continue
        room = limit - used
        if room <= 0:
            cut = True
            break
        if len(line) > room - 1:
            line, cut = line[:max(0, room - 1)], True
        lines.append(line)
        used += len(line) + 1
        if cut:
            break
    if cut and noted:
        lines.append(CUT_NOTE.format(n=used))
    return ("\n".join(lines) + ("\n" if lines else ""))[:MAX_TEXT_CHARS]


def _validate_archive(data: bytes) -> None:
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_ARCHIVE_ENTRIES or sum(info.file_size for info in entries) > MAX_ARCHIVE_BYTES:
            raise UploadTooLarge("Office 文件解压后过大，请拆分后上传")
        if any(info.flag_bits & 1 for info in entries):
            raise UploadError("暂不支持加密的 Office 文件")


def _docx_text(data: bytes) -> str:
    _validate_archive(data)
    with ZipFile(BytesIO(data)) as archive:
        document = ElementTree.fromstring(archive.read("word/document.xml"))
        numbering = (ElementTree.fromstring(archive.read("word/numbering.xml"))
                     if "word/numbering.xml" in archive.namelist() else None)   # clause numbers Word generates
    return docx_document_text(document, MAX_TEXT_CHARS, numbering)


def _xlsx_text(data: bytes) -> str:
    _validate_archive(data)
    try:
        import openpyxl
    except ImportError as exc:
        raise UploadError("Excel 文本解析依赖未安装：请安装 openpyxl") from exc
    workbook = openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=True, keep_links=False)
    try:
        lines: list[str] = []
        used = 0
        has_rows = False
        for sheet in workbook.worksheets:
            heading = f"# {sheet.title}"
            lines.append(heading)
            used += len(heading) + 1
            # Bound sparse or malformed sheets as well as the output text.
            rows = sheet.iter_rows(max_row=min(sheet.max_row or 0, 20_000), max_col=min(sheet.max_column or 0, 256), values_only=True)
            table = table_markdown(rows, max(0, MAX_TEXT_CHARS - used))
            if table:
                has_rows = True
                lines.append(table)
                used += len(table) + 1
            if used >= MAX_TEXT_CHARS:
                break
        return "\n".join(lines) if has_rows else ""
    finally:
        workbook.close()


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise UploadError("PDF 文本解析依赖未安装：请安装 pypdf；扫描件需先完成 OCR") from exc
    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise UploadError("暂不支持加密 PDF，请先解密后上传")
    from packing_assistant.tools.pdf_layout import pages_text

    pages: list[str] = []
    used = 0
    for index, page in enumerate(reader.pages):
        if index >= 400:
            break
        text = page.extract_text() or ""
        pages.append(text)
        used += len(text) + 1
        if used >= MAX_TEXT_CHARS * 2:
            break
    if not any(text.strip() for text in pages):
        return ""
    # paragraphs joined, tables rebuilt, every page under its marker - the same reading a job-folder PDF gets
    return pages_text(pages)


def extract_upload(filename: str, data: bytes) -> tuple[str, str, str]:
    if len(data) > MAX_BYTES:
        raise UploadTooLarge("单文件不能超过 20 MB")
    name = safe_filename(filename)
    kind = Path(name).suffix.lower().lstrip(".")
    if kind not in ALLOWED_EXT:
        raise UploadError("只接受 pdf / docx / xlsx / txt / md / csv / json / log")
    try:
        if kind == "pdf":
            raw, engine = _pdf_text(data), "pypdf"
        elif kind == "docx":
            raw, engine = _docx_text(data), "builtin-docx"
        elif kind == "xlsx":
            raw, engine = _xlsx_text(data), "openpyxl"
        else:
            try:
                raw = data.decode("utf-8-sig")
            except UnicodeDecodeError:
                raw = data.decode("gb18030")
            if "\x00" in raw:
                raise UploadError("文件包含二进制内容，请上传有效的文本文件")
            if kind == "csv":
                raw = csv_text(raw, MAX_TEXT_CHARS)
            engine = "builtin-text"
    except UploadTooLarge:
        raise
    except UploadError as exc:
        # encrypted, binary where text was promised, a parser that is not installed: our own wording
        raise UploadUnreadable(str(exc), name=name, kind=kind, size=len(data)) from exc
    except Exception as exc:
        # Parser messages can echo document bytes; keep those out of responses.
        raise UploadUnreadable(f"{kind} 文件无法解析，请检查文件是否损坏或格式与扩展名一致",
                               name=name, kind=kind, size=len(data)) from exc
    text = _collapse(raw)
    if len(text.strip()) < 8:
        if kind == "pdf":
            message = "PDF 里抽不出可用文字。扫描件需要先 OCR，或另存为 Word/文本"
        elif kind in {"docx", "xlsx"}:
            message = f"{kind} 文件里几乎没有文字内容（不足 8 个字符），请检查是否为空白文档"
        else:
            message = "文件内容几乎为空（不足 8 个字符），请检查后重新上传"
        raise UploadUnreadable(message, name=name, kind=kind, size=len(data))
    return kind, text, engine


def _metadata(directory: Path, identifier: str) -> dict | None:
    if not _ID_RE.fullmatch(identifier):
        return None
    try:
        meta_path = assert_open(_path(directory, f"{identifier}.json"))
        if not meta_path.is_file() or meta_path.stat().st_size > 8192:
            return None
        value = json.loads(meta_path.read_text(encoding="utf-8"))
        if not isinstance(value, dict) or value.get("id") != identifier:
            return None
        if not all(isinstance(value.get(key), str) for key in ("name", "kind", "parse")):
            return None
        if value["kind"] not in ALLOWED_EXT or not all(
            isinstance(value.get(key), int) and value[key] >= 0 for key in ("bytes", "chars")
        ):
            return None
        if not all(assert_open(_path(directory, f"{identifier}.{ext}")).is_file() for ext in ("bin", "txt")):
            return None
        return {key: value[key] for key in ("id", "name", "kind", "bytes", "chars", "parse")}
    except (OSError, ValueError, RuntimeError):
        return None


def list_uploads(session: str) -> list[dict]:
    directory = assert_open(_directory(session))
    if not directory.is_dir():
        return []
    result = []
    for path in directory.glob("*.json"):
        meta = _metadata(directory, path.stem)
        if meta is not None:
            result.append(meta)
    return sorted(result, key=lambda item: item["name"])


def _unreadable_lines(path: Path) -> list[dict]:
    try:
        log = assert_open(path)
        if not log.is_file() or log.stat().st_size > 64_000:
            return []
        rows = [json.loads(line) for line in log.read_text(encoding="utf-8").splitlines() if line.strip()]
    except (OSError, ValueError, RuntimeError):
        return []
    return [row for row in rows if isinstance(row, dict) and isinstance(row.get("name"), str)
            and isinstance(row.get("reason"), str)][-MAX_UNREADABLE:]


def _remember_unreadable(session: str, failures: list[UploadUnreadable]) -> None:
    """Best effort. The upload is refused either way; the note only lets a later check tell a file that
    was given and not read from a file that was never given."""
    try:
        with _LOCK:
            path = _path(_directory(session), UNREADABLE_LOG)
            rows = _unreadable_lines(path) + [
                {"name": exc.name, "kind": exc.kind, "bytes": exc.size, "reason": str(exc)[:200], "at": int(time.time())}
                for exc in failures]
            guarded_write_text(path, "".join(json.dumps(row, ensure_ascii=False) + "\n" for row in rows[-MAX_UNREADABLE:]))
    except (OSError, ValueError, RuntimeError):
        pass


def unreadable_uploads(session: str) -> list[dict]:
    """Files this task tried to attach that gave no text: the last attempt per name, without the names
    that have been attached successfully since (the same file after OCR, say)."""
    with _LOCK:
        directory = assert_open(_directory(session))
        if not directory.is_dir():
            return []
        attached = {item["name"] for item in list_uploads(session)}
        latest: dict[str, dict] = {}
        for row in _unreadable_lines(_path(directory, UNREADABLE_LOG)):
            latest[row["name"]] = row
    return [{"name": name, "kind": str(row.get("kind") or ""), "reason": row["reason"]}
            for name, row in latest.items() if name not in attached]


def save_uploads(session: str, files: Iterable[tuple[str, bytes]]) -> dict:
    directory = _directory(session)
    files = list(files)
    if not files:
        raise UploadError("没有收到文件")
    if len(files) > MAX_FILES:
        raise UploadError("同一会话最多 12 个附件")
    if sum(len(data) for _, data in files) > MAX_REQUEST_BYTES:
        raise UploadTooLarge("一次上传不能超过 25 MB")
    prepared = []
    unreadable: list[UploadUnreadable] = []
    for name, data in files:
        try:
            prepared.append((safe_filename(name), data, extract_upload(name, data)))
        except UploadUnreadable as exc:
            unreadable.append(exc)
    if unreadable:
        # still nothing of this batch is saved; what changes is that the refusal leaves a note
        _remember_unreadable(session, unreadable)
        raise unreadable[0]
    saved: list[dict] = []
    created: list[Path] = []
    with _LOCK:
        if len(list_uploads(session)) + len(prepared) > MAX_FILES:
            raise UploadError("同一会话最多 12 个附件")
        try:
            for name, data, (kind, text, engine) in prepared:
                identifier = uuid4().hex[:12]
                while any((directory / f"{identifier}.{ext}").exists() for ext in ("bin", "txt", "json")):
                    identifier = uuid4().hex[:12]
                meta = {"id": identifier, "name": name, "kind": kind, "bytes": len(data), "chars": len(text), "parse": engine}
                for extension, content in (("bin", data), ("txt", text.encode("utf-8"))):
                    path = _path(directory, f"{identifier}.{extension}")
                    created.append(assert_write(path))
                    guarded_write_bytes(path, content)
                pending = _path(directory, f"{identifier}.json.tmp")
                created.append(assert_write(pending))
                guarded_write_text(pending, json.dumps(meta, ensure_ascii=False))
                published = assert_write(_path(directory, f"{identifier}.json"))
                pending.replace(published)
                created.append(published)
                saved.append(meta)
        except Exception:
            for path in reversed(created):
                try:
                    assert_write(path).unlink(missing_ok=True)
                except OSError:
                    pass
            raise
    return {"ok": True, "files": saved}


def save_upload(session: str, filename: str, data: bytes) -> dict:
    return save_uploads(session, [(filename, data)])["files"][0]


def read_upload(session: str, identifier: str, offset: int = 0, limit: int = 8000) -> str:
    if not isinstance(identifier, str) or not _ID_RE.fullmatch(identifier):
        raise UploadError("附件 id 无效")
    if not isinstance(offset, int) or offset < 0 or not isinstance(limit, int) or limit < 0:
        raise UploadError("附件读取 offset 和 limit 必须为非负整数")
    directory = _directory(session)
    meta = _metadata(directory, identifier)
    if meta is None:
        raise UploadError("附件不存在")
    take = min(limit or 8000, 20_000)
    with assert_open(_path(directory, f"{identifier}.txt")).open(encoding="utf-8") as stream:
        text = stream.read(MAX_TEXT_CHARS)
    selected = text[offset:offset + take]
    more = max(0, len(text) - offset - len(selected))
    return f"【用户上传：{meta['name']}】offset={offset} 本段{len(selected)}字 剩余约{more}字\n\n{selected}"


def extracted_documents(session: str, ids: Iterable[str]) -> list[dict]:
    """Read selected local extraction caches for indexing, without prefix clipping."""
    directory = _directory(session)
    result = []
    for identifier in dict.fromkeys(ids):
        meta = _metadata(directory, identifier)
        if meta is None:
            raise UploadError("附件不存在，请重新选择当前任务的附件")
        with assert_open(_path(directory, f"{identifier}.txt")).open(encoding="utf-8") as stream:
            text = stream.read(MAX_TEXT_CHARS)
        result.append({**meta, "text": text})
    return result


def strict_documents(session: str) -> list[dict]:
    """Complete attachment snapshot for repair, never silently omit bad records.

    Ordinary browsing can tolerate a damaged upload entry. A destructive refresh
    of derived indexes must also notice orphaned original/text files, so their
    last usable search entries are not discarded as if they had been removed.
    """
    with _LOCK:
        directory = assert_open(_directory(session))
        if not directory.exists():
            return []
        if not directory.is_dir():
            raise UploadError("附件目录不可读，请检查原始资料后重建")
        identifiers = set()
        for path in directory.iterdir():
            if path.suffix.lower() not in {".json", ".txt", ".bin"}:
                continue
            if not _ID_RE.fullmatch(path.stem):
                raise UploadError("附件目录含无效记录，请检查或重新上传后重建")
            identifiers.add(path.stem)
        if len(identifiers) > MAX_FILES:
            raise UploadError("附件记录超过任务上限，请核对后重建")
        result = []
        for identifier in sorted(identifiers):
            meta = _metadata(directory, identifier)
            if (meta is None or type(meta["bytes"]) is not int or type(meta["chars"]) is not int
                    or not 0 <= meta["bytes"] <= MAX_BYTES or not 0 <= meta["chars"] <= MAX_TEXT_CHARS):
                raise UploadError("部分附件记录损坏或文件缺失，请检查或重新上传后重建")
            original = assert_open(_path(directory, identifier + ".bin"))
            if original.stat().st_size != meta["bytes"]:
                raise UploadError("附件原件大小与记录不一致，请核对原件后重建")
            try:
                with assert_open(_path(directory, identifier + ".txt")).open(encoding="utf-8", newline="") as stream:
                    text = stream.read(MAX_TEXT_CHARS + 1)
            except UnicodeError as exc:
                raise UploadError("附件提取文本损坏，请重新上传原件后重建") from exc
            if len(text) != meta["chars"]:
                raise UploadError("附件提取文本不完整，请核对或重新上传原件后重建")
            result.append({**meta, "text": text})
        return result


def attachment_context(session: str, ids: Iterable[str] | None = None) -> str:
    """Only explicitly selected attachments enter chat; an empty list adds none."""
    identifiers = list(ids or [])
    if not identifiers:
        return ""
    if len(identifiers) > MAX_FILES:
        raise UploadError("同一会话最多 12 个附件")
    parts: list[str] = []
    used = 0
    for identifier in dict.fromkeys(identifiers):
        content = read_upload(session, identifier, limit=20_000)
        room = INJECT_CHARS - used
        if room <= 0:
            break
        parts.append(content[:room])
        used += len(parts[-1]) + 2
    return "\n\n".join(parts)[:INJECT_CHARS]


def bundle_for_prompt(session: str, ids: Iterable[str], user_message: str) -> str:
    context = attachment_context(session, ids)
    if not context:
        return user_message
    return f"{context}\n\n---\n用户说：\n{user_message}"
