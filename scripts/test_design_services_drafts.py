#!/usr/bin/env python3
"""T045 professional sections, scoped facts, truthful limits and actual Office outputs."""

from __future__ import annotations

import os
from pathlib import Path
import re
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from packing_assistant.expert_roster import get_expert
from packing_assistant.office_job import tables_from_md
from packing_assistant.post_drafts.design_services import MISSING, build_draft
from packing_assistant.tools.tender_review import forbidden_hits

CASES = {
    "plumbing": {"tool": "plumbing__memo", "title": "排水与标高衔接", "field": "出户标高", "value": "1.20m",
                 "text": "系统：生活排水；室内地面标高：1.50m；出户标高：1.20m；市政井标高：1.00m；流量：3.1L/s；管径：用户原表DN100",
                 "chapters": ("水源与市政接驳", "生活给水与分区", "雨水、溢流与回用", "消防水量与泵房资料", "特殊部位", "节水与计量")},
    "hvac": {"tool": "hvac__memo", "title": "冷热源方案与负荷资料", "field": "用户冷负荷", "value": "150kW",
             "text": "系统：办公空调；冷负荷：150kW；热负荷：90kW；新风量：1800m³/h；室内参数：用户参数表A；冷热源：用户拟议冷水机组",
             "chapters": ("系统范围与空气计算参数", "风系统与水系统", "防烟、排烟与联锁", "机房、冷却塔与竖向空间", "消声、隔振、冷凝水与保温", "节能资料")},
    "electrical": {"tool": "electrical__memo", "title": "负荷计算书数据登记", "field": "设备容量", "value": "120kW",
                   "text": "系统：办公配电；设备容量：120kW；需要系数：0.7；供电电压：400V；变压器容量：用户原表160kVA；接地电阻目标：用户文件2Ω",
                   "chapters": ("系统范围与市政电源", "变配电方案核对", "照明与疏散指示", "防雷、接地与等电位", "线路敷设", "电气消防接口", "弱电电源与接地交接")},
    "fire-protect": {"tool": "fire-protect__brief", "title": "安全疏散与避难", "field": "用户疏散宽度", "value": "1.40m",
                     "text": "单体：办公楼A；建筑分类：用户待复核分类；建筑高度：45m；层数：12层；安全出口：用户图示2处；疏散宽度：1.40m；防火分区：分区A",
                     "chapters": ("工程概况", "总图与消防救援条件", "建筑防火与分区", "消防给水与灭火", "防烟排烟", "火灾报警与联动", "消防电气", "装修与既有使用边界", "报审资料待办目录")},
    "steel": {"tool": "steel__memo", "title": "材料与构件规格记录", "field": "用户既有/提出规格", "value": "用户原表H300×150",
              "text": "构件：梁G1；结构体系：用户拟议框架；跨度：12m；钢材牌号：Q355B；用户规格：用户原表H300×150；风压：0.55kPa；连接方式：用户节点焊接",
              "chapters": ("结构体系与构件范围", "荷载与计算资料", "连接与节点资料", "稳定与支撑", "防腐涂装", "构件防火保护", "加工、运输、安装与检测资料", "实体连接界面")},
}


def draft(post: str, text: str = "写一份内部讨论草稿") -> str:
    result = build_draft(post, CASES[post]["tool"], text)
    assert isinstance(result, str)
    return result


def rows_for(markdown: str, heading: str) -> list[list[str]]:
    return next(rows for name, rows in tables_from_md(markdown) if name.startswith(heading))


def field(rows: list[list[str]], name: str, row: int = 1) -> str:
    return rows[row][rows[0].index(name)]


class DesignServicesBuilderTests(unittest.TestCase):
    def test_exact_roster_post_tool_contract(self) -> None:
        for post, case in CASES.items():
            self.assertEqual((case["tool"],), get_expert(post).exclusive)
            self.assertIsNotNone(build_draft(post, case["tool"], ""))
            self.assertIsNone(build_draft(post, "different__tool", ""))
        self.assertIsNone(build_draft("architecture", "architecture__memo", ""))

    def test_each_post_has_its_own_professional_sections(self) -> None:
        for post, case in CASES.items():
            with self.subTest(post=post):
                md = draft(post)
                for heading in (case["title"], *case["chapters"]):
                    self.assertIn("## " + heading, md)
                self.assertNotIn("## 用户原文", md)
                self.assertIn("草稿声明", md)
                self.assertIn("submit_blocked=true", md)
                self.assertFalse(forbidden_hits(md))
        self.assertNotIn("防腐涂装", draft("electrical"))
        self.assertNotIn("用户变压器容量", draft("plumbing"))

    def test_unknown_parameters_never_become_sizes_or_calculations(self) -> None:
        for post in CASES:
            with self.subTest(post=post):
                md = draft(post)
                self.assertIn("- 辖区：UNSPECIFIED\n", md)
                self.assertIn(MISSING, md)
                self.assertNotIn("见图", md)
                self.assertNotIn("经验算满足", md)
                self.assertNotRegex(md, r"\d+(?:\.\d+)?\s*(?:kW|kVA|MPa|kPa|mm|L/s|小时|米)\b")
                for _, rows in tables_from_md(md):
                    for index, heading in enumerate(rows[0]):
                        if heading.startswith("本稿"):
                            self.assertTrue(all(row[index] == MISSING for row in rows[1:]))

    def test_explicit_values_enter_their_professional_table(self) -> None:
        for post, case in CASES.items():
            with self.subTest(post=post):
                md = draft(post, case["text"] + "；资料来源：用户记录A；图号：D-01；计算书：用户计算资料A")
                self.assertEqual(case["value"], field(rows_for(md, case["title"]), case["field"]))
                self.assertEqual("用户记录A", field(rows_for(md, case["title"]), "资料来源"))
                self.assertIn("用户参数仅抄录待核", md)
                self.assertNotIn("见图 D-01", md)
                for _, rows in tables_from_md(md):
                    for index, heading in enumerate(rows[0]):
                        if heading.startswith("本稿"):
                            self.assertTrue(all(row[index] == MISSING for row in rows[1:]))

    def test_column_order_and_empty_values_do_not_cross_system_or_component(self) -> None:
        examples = (
            ("plumbing", "水压|系统|出户标高", "0.28MPa|给水A|1.20m", "|给水B|", "水源与市政接驳", "用户水压", "0.28MPa"),
            ("hvac", "冷负荷|系统|新风量", "150kW|空调A|1800m³/h", "|空调B|", "冷热源方案与负荷资料", "用户冷负荷", "150kW"),
            ("electrical", "设备容量|系统|需要系数", "120kW|配电A|0.7", "|配电B|", "负荷计算书数据登记", "设备容量", "120kW"),
            ("fire-protect", "疏散宽度|单体|安全出口", "1.40m|办公楼A|图示2处", "|办公楼B|", "安全疏散与避难", "用户疏散宽度", "1.40m"),
            ("steel", "用户规格|构件|跨度", "原表H300|梁A|12m", "|梁B|", "材料与构件规格记录", "用户既有/提出规格", "原表H300"),
        )
        for post, headers, one, two, title, name, value in examples:
            with self.subTest(post=post):
                source = f"|{headers}|\n|---|---|---|\n|{one}|\n|{two}|"
                rows = rows_for(draft(post, source), title)
                self.assertEqual(3, len(rows))
                self.assertEqual(value, field(rows, name, 1))
                self.assertEqual(MISSING, field(rows, name, 2))

    def test_missing_identity_stays_its_own_table_row(self) -> None:
        rows = rows_for(draft("steel", "| 构件 | 用户规格 |\n| --- | --- |\n| 梁A | 原表H300 |\n| | 原表H400 |"), "材料与构件规格记录")
        self.assertEqual(MISSING, field(rows, "构件", 2))
        self.assertEqual("原表H400", field(rows, "用户既有/提出规格", 2))
        self.assertEqual("原表H300", field(rows, "用户既有/提出规格", 1))

    def test_multiline_labels_preserve_record_even_when_number_precedes_name(self) -> None:
        rows = rows_for(draft("hvac", "项目：示例；单体：办公楼；辖区：CN\n冷负荷：150kW；系统：空调A\n冷负荷：90kW；系统：空调B"), "冷热源方案与负荷资料")
        self.assertEqual(["空调A", "空调B"], [row[2] for row in rows[1:]])
        self.assertEqual(["150kW", "90kW"], [field(rows, "用户冷负荷", index) for index in (1, 2)])
        self.assertEqual(["办公楼", "办公楼"], [row[1] for row in rows[1:]])

    def test_key_value_table_and_chapter_parameter_semantics(self) -> None:
        md = draft("plumbing", "| 字段 | 内容 |\n| --- | --- |\n| 系统 | 排水A |\n| 室内地面标高 | 1.50m |\n| 出户标高 | 1.20m |\n| 市政井标高 | 1.00m |")
        rows = rows_for(md, "排水与标高衔接")
        self.assertEqual("排水A", field(rows, "系统"))
        self.assertEqual("1.20m", field(rows, "出户标高"))
        self.assertEqual(MISSING, field(rows, "本稿管径计算结果"))

    def test_dual_sources_and_interfaces_are_separate(self) -> None:
        for post in CASES:
            with self.subTest(post=post):
                md = draft(post, "辖区：DUAL\nCN依据：CN用户文件；SG依据：SG用户文件\nCN接口：CN接口清单；SG接口：SG接口清单")
                cn = rows_for(md, "依据与接口 · CN")
                sg = rows_for(md, "依据与接口 · SG")
                self.assertEqual("CN用户文件", field(cn, "用户依据文件"))
                self.assertEqual("SG用户文件", field(sg, "用户依据文件"))
                self.assertEqual("CN接口清单", field(cn, "用户接口要求"))
                self.assertNotIn("SG用户文件", str(cn))
                self.assertNotIn("CN用户文件", str(sg))
                self.assertIn("- 辖区：DUAL\n", md)

    def test_explicit_record_jurisdictions_and_unassigned_dual_material(self) -> None:
        md = draft("plumbing", "辖区：CN；系统：给水CN；水压：0.28MPa；依据：CN文件\n辖区：SG；系统：给水SG；依据：SG文件")
        rows = rows_for(md, "水源与市政接驳")
        self.assertEqual(["CN", "SG"], [row[0] for row in rows[1:]])
        self.assertEqual(MISSING, field(rows, "用户水压", 2))
        self.assertEqual("CN文件", field(rows_for(md, "依据与接口 · CN"), "用户依据文件"))
        unknown = draft("hvac", "辖区：DUAL；依据：待分配文件；系统：空调")
        self.assertIn("依据与接口 · 辖区 A（待指明）", unknown)
        self.assertIn("依据与接口 · 辖区 B（待指明）", unknown)
        self.assertIn("跨辖区资料待分配", unknown)
        self.assertNotIn("依据与接口 · CN", unknown)

    def test_jurisdictions_do_not_spill_between_objects(self) -> None:
        for post, case in CASES.items():
            identity = "构件" if post == "steel" else "系统"
            for source in (
                f"{identity}：甲；辖区：CN\n{identity}：乙\n{identity}：丙；辖区：SG",
                f"辖区：CN；{identity}：甲\n{identity}：乙\n辖区：SG；{identity}：丙",
                f"|{identity}|辖区|\n|---|---|\n|甲|CN|\n|乙||\n|丙|SG|",
                f"|字段|内容|\n|---|---|\n|{identity}|甲|\n|辖区|CN|\n|{identity}|乙|\n|{identity}|丙|\n|辖区|SG|",
            ):
                with self.subTest(post=post, source=source):
                    rows = rows_for(draft(post, source), case["title"])
                    self.assertEqual(["甲", "乙", "丙"], [row[2] for row in rows[1:]])
                    self.assertEqual(["CN", "UNSPECIFIED", "SG"], [row[0] for row in rows[1:]])

    def test_only_top_level_region_is_a_default_and_explicit_empty_overrides_it(self) -> None:
        for post, case in CASES.items():
            identity = "构件" if post == "steel" else "系统"
            for header in ("辖区：EU", "欧盟项目", "|字段|内容|\n|---|---|\n|辖区|EU|"):
                with self.subTest(post=post, header=header):
                    source = f"{header}\n{identity}：甲\n辖区：CN\n{identity}：乙\n{identity}：丙；辖区："
                    rows = rows_for(draft(post, source), case["title"])
                    self.assertEqual(["甲", "乙", "丙"], [row[2] for row in rows[1:]])
                    self.assertEqual(["CN", "EU", "UNSPECIFIED"], [row[0] for row in rows[1:]])

    def test_unknown_object_basis_is_not_assigned_to_a_known_sibling_jurisdiction(self) -> None:
        md = draft("plumbing", "系统：甲；辖区：CN；依据：甲文件\n系统：乙；依据：乙未知辖区文件")
        cn = rows_for(md, "依据与接口 · CN")
        self.assertIn("甲文件", str(cn))
        self.assertNotIn("乙未知辖区文件", str(cn))
        self.assertIn("乙未知辖区文件", str(rows_for(md, "跨辖区资料待分配")))

    def test_escaping_is_only_at_render_time(self) -> None:
        raw = "A&B < 条件 | 保留实体 &lt;"
        for post, case in CASES.items():
            md = draft(post, case["text"] + "；资料来源：" + raw)
            self.assertEqual(raw, field(rows_for(md, case["title"]), "资料来源"))
            self.assertIn("A&amp;B &lt; 条件 &#124; 保留实体 &amp;lt;", md)
            self.assertNotIn("A&amp;amp;B", md)


class DesignServicesRuntimeTests(unittest.TestCase):
    def setUp(self) -> None:
        from packing_assistant import expert_turn
        from packing_assistant.runtime.scheduler import Scheduler

        temp = tempfile.TemporaryDirectory(prefix="civil-design-services-")
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        self.posts = expert_turn
        for context in (
            patch.object(expert_turn, "_OUT", self.root / "sessions"),
            patch.dict(os.environ, {"CIVIL_SANDBOX_ROOTS": str(self.root), "CIVIL_JOB_ROOT": "", "CIVIL_SANDBOX": "workspace-write", "CIVIL_APPROVAL": "on-request", "PYTHON_DOTENV_DISABLED": "1"}),
            patch("packing_assistant.runtime.memory.assemble_context", return_value={}),
            patch("packing_assistant.runtime.scheduler.get_scheduler", return_value=Scheduler()),
        ):
            context.start()
            self.addCleanup(context.stop)

    def test_each_post_writes_actual_markdown_and_xlsx_with_original_values(self) -> None:
        import openpyxl

        raw = "原始 < 参数 | A&B &lt;"
        for post, case in CASES.items():
            with self.subTest(post=post):
                result = self.posts.run_expert_turn(case["text"] + "；资料来源：" + raw, post, force_intent="run", confirm_ok=True, session_id=post)
                self.assertTrue(result["ok"], result)
                self.assertTrue(result["wrote"], result)
                self.assertEqual("done", result["state"])
                files = [Path(value["path"]) for value in result["files"]]
                md = next(path for path in files if path.suffix == ".md")
                xlsx = next(path for path in files if path.suffix == ".xlsx")
                self.assertTrue(md.is_relative_to(self.root))
                self.assertTrue(xlsx.is_relative_to(self.root))
                self.assertEqual(case["value"], field(rows_for(md.read_text(encoding="utf-8"), case["title"]), case["field"]))
                workbook = openpyxl.load_workbook(xlsx)
                try:
                    sheet = next(workbook[name] for name in workbook.sheetnames if name.startswith(case["title"]))
                    values = list(sheet.iter_rows(values_only=True))
                    self.assertIn(case["value"], values[1])
                    self.assertIn(raw, values[1])
                finally:
                    workbook.close()

    def test_questions_do_not_write_any_post(self) -> None:
        for post in CASES:
            result = self.posts.run_expert_turn("需要准备哪些设计资料？", post, force_intent="chat", session_id="question-" + post)
            self.assertTrue(result["ok"], result)
            self.assertFalse(result["wrote"])
            self.assertEqual([], result["files"])
        self.assertFalse(list(self.root.rglob("*.md")))

    def test_multirecord_excel_keeps_missing_quantities_in_the_correct_row(self) -> None:
        import openpyxl

        for post, case in CASES.items():
            with self.subTest(post=post):
                identity = "构件" if post == "steel" else "系统"
                input_field = {"plumbing": "出户标高", "hvac": "冷负荷", "electrical": "设备容量",
                               "fire-protect": "疏散宽度", "steel": "用户规格"}[post]
                source = f"|{input_field}|{identity}|\n|---|---|\n|{case['value']}|记录甲|\n||记录乙|"
                result = self.posts.run_expert_turn(source, post, force_intent="run", confirm_ok=True, session_id="multi-" + post)
                self.assertTrue(result["ok"], result)
                files = [Path(item["path"]) for item in result["files"]]
                content = next(path for path in files if path.suffix == ".md").read_text(encoding="utf-8")
                rows = rows_for(content, case["title"])
                self.assertEqual(3, len(rows))
                self.assertEqual(MISSING, field(rows, case["field"], 2))
                workbook = openpyxl.load_workbook(next(path for path in files if path.suffix == ".xlsx"))
                try:
                    sheet = next(workbook[name] for name in workbook.sheetnames if name.startswith(case["title"]))
                    exported = list(sheet.iter_rows(values_only=True))
                    index = exported[0].index(case["field"])
                    self.assertEqual(case["value"], exported[1][index])
                    self.assertEqual(MISSING, exported[2][index])
                    self.assertIn("记录乙", exported[2])
                finally:
                    workbook.close()

    def test_high_risk_requires_a_real_confirmation(self) -> None:
        from packing_assistant.runtime.tool_engine import default_engine

        for post in ("fire-protect", "steel"):
            self.assertEqual("high", get_expert(post).risk)
            for confirmation in (False, "false"):
                result = default_engine().execute(CASES[post]["tool"], {"text": "写内部草稿", "confirm_ok": confirmation, "session_id": "gate-" + post}, expert_id=post)
                self.assertFalse(result.get("wrote"), result)
                if confirmation is False:
                    self.assertTrue(result.get("hitl_pending"), result)
                else:
                    self.assertEqual(result.get("error_code"), "invalid_args", result)
                self.assertFalse(result.get("files"))
        self.assertFalse(list(self.root.rglob("*.md")))

    def test_export_keeps_object_jurisdictions_independent(self) -> None:
        import openpyxl

        for post, case in CASES.items():
            with self.subTest(post=post):
                identity = "构件" if post == "steel" else "系统"
                source = f"{identity}：甲；辖区：CN\n{identity}：乙\n{identity}：丙；辖区：SG"
                result = self.posts.run_expert_turn(source, post, force_intent="run", confirm_ok=True, session_id="zones-" + post)
                self.assertTrue(result["ok"], result)
                files = [Path(item["path"]) for item in result["files"]]
                content = next(path for path in files if path.suffix == ".md").read_text(encoding="utf-8")
                self.assertEqual(["CN", "UNSPECIFIED", "SG"], [row[0] for row in rows_for(content, case["title"])[1:]])
                workbook = openpyxl.load_workbook(next(path for path in files if path.suffix == ".xlsx"))
                try:
                    sheet = next(workbook[name] for name in workbook.sheetnames if name.startswith(case["title"]))
                    exported = list(sheet.iter_rows(values_only=True))
                    self.assertEqual(["CN", "UNSPECIFIED", "SG"], [row[0] for row in exported[1:]])
                    self.assertEqual(["甲", "乙", "丙"], [row[2] for row in exported[1:]])
                finally:
                    workbook.close()

    def test_read_only_blocks_direct_and_tool_entry_for_every_post(self) -> None:
        from packing_assistant.runtime.tool_engine import default_engine

        with patch.dict(os.environ, {"CIVIL_SANDBOX": "read-only"}):
            for post, case in CASES.items():
                direct = self.posts.run_expert_turn("写内部草稿", post, force_intent="run", confirm_ok=True, session_id="readonly-" + post)
                tool = default_engine().execute(case["tool"], {"text": "写内部草稿", "confirm_ok": True, "session_id": "readonly-tool-" + post}, expert_id=post)
                for result in (direct, tool):
                    self.assertFalse(result["ok"], result)
                    self.assertFalse(result.get("wrote"))
                    self.assertFalse(result.get("files"))
        self.assertFalse(list(self.root.rglob("*.md")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
