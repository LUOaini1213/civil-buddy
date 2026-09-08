#!/usr/bin/env python3
"""对着网上公开的真实装箱单 / packing list 测解析 + demo 装箱。

文件缓存在 data/external/web_packing_lists/（已下载则不再拉）。
每个样本两步：
  1) 原生 parse_table_file（第一行当表头，xlsx only）—— 真实箱单几乎都会失败
  2) 箱单适配：扫表头行、拆 LxWxH 字符串、件数优先用 cartons/packages
有外尺寸才跑 Team A/B，对照文档上的柜数/件数/重量。
"""

from __future__ import annotations

import json
import re
import sys
import time
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence, Tuple

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

CACHE = ROOT / "data" / "external" / "web_packing_lists"
OUT_DIR = ROOT / "output" / "web_packing_lists"

DIM_RE = re.compile(
    r"(?P<l>\d+(?:\.\d+)?)\s*[xX×*]\s*(?P<w>\d+(?:\.\d+)?)\s*[xX×*]\s*(?P<h>\d+(?:\.\d+)?)"
    r"\s*(?P<u>cm|mm|m|in|inch|inches)?"
    r"(?:\s*(?P<u2>cm|mm|m|in|inch|inches))?",
    re.I,
)

# 公开可下载的真实箱单 / 带样例数据的 packing list（不是我们自己合成的）
SOURCES: List[Dict[str, Any]] = [
    {
        "id": "dpx_sample_pl",
        "file": "dpx_sample_ci_pl.xls",
        "url": "https://www.dpxlogistics.com/wp-content/uploads/DPX-Sample-Commercial-Invoice-and-Packing-List-1.xls",
        "kind": "filled_packing_list",
        "sheet": "Sample Packing List",
        "doc_truth": {
            "packages": 2,
            "gross_kg": 8.0,
            "pieces": [
                {"name": "3 sets of Book (Company Report)", "lwh_cm": [20, 30, 25], "gw_kg": 4.5},
                {
                    "name": "1 set of Book (Company Report) and 4 CDs",
                    "lwh_cm": [15, 15, 30],
                    "gw_kg": 3.5,
                },
            ],
            "mode": "air_small",
            "expected_containers": 1,
            "note": "DPX 2014 泰→新空运样例：2 纸箱，合并尺寸字段 20x30x25 cm",
        },
    },
    {
        "id": "supplyautomate_pl",
        "file": "supplyautomate_packing_list.xlsx",
        "url": "https://www.supplyautomate.com/templates/packing-list-template.xlsx",
        "kind": "filled_packing_list",
        "sheet": "Packing List",
        "doc_truth": {
            "packages": 40,
            "carton_lwh_cm": [60, 40, 35],
            "gross_kg_per_carton": 12.4,
            "gross_kg": 496.0,
            "cbm": 3.36,
            "container_no": "MSCU1234567",
            "expected_containers": 1,
            "route": "Yantian → Los Angeles",
            "note": "模板自带填好的第 15 行：40 箱帆布袋，一柜号",
        },
    },
    {
        "id": "skydo_export_pl",
        "file": "skydo_packing_list_template.xlsx",
        "url": "https://skydo-assets.s3.ap-south-1.amazonaws.com/strapi-assets/packing_list_template_b8bae18708.xlsx",
        "kind": "filled_packing_list_no_dims",
        "sheet": "Packing List",
        "doc_truth": {
            "packages": 24,
            "units": 1000,
            "gross_kg": 368.0,
            "net_kg": 327.5,
            "mode": "Sea (FCL)",
            "container_no": "MSCU 8874920",
            "expected_containers": 1,
            "has_lwh": False,
            "note": "印度成衣 FCL 样例：24 件/368 kg，无纸箱外尺寸",
        },
    },
    {
        "id": "hansatic_template",
        "file": "hansatic_packing_list_template.xlsx",
        "url": "https://hansatic.com/templates/hansatic-packing-list-template.xlsx",
        "kind": "blank_template",
        "sheet": "Packing List",
        "doc_truth": {
            "packages": 0,
            "note": "货代级空白模板：有 L/W/H cm 列，无实货",
        },
    },
    {
        "id": "incodocs_pdf",
        "file": "incodocs_packing_list.pdf",
        "url": "https://incodocs.com/templates/pdfs/Packing%20List.pdf",
        "kind": "blank_template",
        "doc_truth": {
            "packages": 0,
            "note": "IncoDocs 空白 PDF 模板，无行项目",
        },
    },
    {
        "id": "piston_widget_example",
        "file": "piston_packaging_workbook.xlsx",
        "url": "https://www.pistonautomotive.com/wp-content/uploads/2023/08/Piston-Automotive-Packaging-Workbook.xlsx",
        "kind": "packaging_spec_example",
        "sheet": "Production Packaging Example",
        "doc_truth": {
            "packages": 27,
            "carton_lwh_in": [15, 15, 13],
            "full_container_lb": 34,
            "parts_per_carton": 150,
            "containers_per_pallet": 27,
            "expected_containers": 1,
            "note": "汽车包装承认书样例：15×15×13 in 纸箱，27 箱/托",
        },
    },
    {
        "id": "smartsheet_ci",
        "file": "smartsheet_commercial_invoice.xlsx",
        "url": "https://www.smartsheet.com/sites/default/files/commercial-invoice-template.xlsx",
        "kind": "not_packing_list",
        "doc_truth": {"note": "商业发票模板，不是箱单"},
    },
]


def _download(src: Dict[str, Any]) -> Path:
    CACHE.mkdir(parents=True, exist_ok=True)
    dest = CACHE / src["file"]
    if dest.exists() and dest.stat().st_size > 200:
        return dest
    import urllib.request

    req = urllib.request.Request(
        src["url"],
        headers={"User-Agent": "Mozilla/5.0 (compatible; packing-agent-test/1.0)"},
    )
    with urllib.request.urlopen(req, timeout=45) as resp:
        dest.write_bytes(resp.read())
    return dest


def _sheet_rows(path: Path, sheet: Optional[str] = None) -> Tuple[str, List[List[Any]]]:
    suf = path.suffix.lower()
    if suf == ".xls":
        import xlrd

        wb = xlrd.open_workbook(str(path))
        name = sheet if sheet and sheet in wb.sheet_names() else wb.sheet_names()[0]
        sh = wb.sheet_by_name(name)
        rows = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
        return name, rows
    if suf in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, data_only=True, read_only=True)
        name = sheet if sheet and sheet in wb.sheetnames else wb.sheetnames[0]
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
        return name, rows
    return "", []


def _blob(row: Sequence[Any]) -> str:
    return " ".join(str(c or "") for c in row)


def _norm(h: Any) -> str:
    s = str(h or "").strip().lower()
    s = s.replace("\n", " ").replace("&amp;", "&")
    s = re.sub(r"\s+", " ", s)
    return s


def _is_header_row(row: Sequence[Any]) -> bool:
    blob = _blob(row).lower()
    if len(blob) < 8:
        return False
    has_name = any(
        k in blob
        for k in ("description", "content", "品名", "货物", "sku", "item description")
    )
    has_dim = any(
        k in blob
        for k in (
            "dimension",
            "l (cm)",
            "l cm",
            "length",
            "measurement",
            "长",
            "lwh",
        )
    )
    has_wt = any(k in blob for k in ("weight", "gross", "净重", "毛重", "kg"))
    has_qty = any(
        k in blob for k in ("qty", "quantity", "packages", "cartons", "数量", "箱数", "件数")
    )
    return has_name and (has_dim or (has_wt and has_qty))


def _scale_to_mm(val: float, unit: str) -> float:
    u = (unit or "").lower()
    if u in ("in", "inch", "inches"):
        return val * 25.4
    if u == "m":
        return val * 1000.0
    if u == "mm":
        return val
    if u == "cm":
        return val * 10.0
    return val * 10.0 if val <= 300 else val


def _parse_dim_cell(text: Any) -> Optional[Tuple[float, float, float]]:
    if text is None:
        return None
    s = str(text).strip()
    m = DIM_RE.search(s)
    if not m:
        return None
    unit = m.group("u") or m.group("u2") or ""
    if not unit:
        low = s.lower()
        for cand in ("cm", "mm", "inch", "inches", " in"):
            if cand in low:
                unit = "in" if "in" in cand else cand.strip()
                break
    l, w, h = float(m.group("l")), float(m.group("w")), float(m.group("h"))
    return _scale_to_mm(l, unit), _scale_to_mm(w, unit), _scale_to_mm(h, unit)


def _map_headers(headers: Sequence[Any]) -> Dict[str, int]:
    """箱单列 → 字段下标。件数优先 cartons/packages，避开 qty per carton。"""
    idx: Dict[str, int] = {}
    norms = [_norm(h) for h in headers]
    for i, h in enumerate(norms):
        if not h:
            continue
        if "qty per carton" in h or "quantity per carton" in h or "units per" in h:
            idx.setdefault("inner_qty", i)
            continue
        if h in ("cartons", "packages", "箱数", "件数") or h.startswith("kind & no"):
            idx["quantity"] = i
            continue
        if h in ("qty", "qty.", "q'ty", "quantity", "数量") and "quantity" not in idx:
            idx["quantity"] = i
            continue
        if any(
            k == h or k in h
            for k in (
                "description of goods",
                "goods description",
                "content",
                "description",
                "品名",
                "货物名称",
            )
        ):
            idx.setdefault("name", i)
            continue
        if h in ("sku", "item", "s.no.", "s.no", "package no.", "package no", "carton no."):
            idx.setdefault("id", i)
            continue
        if h in ("l cm", "l (cm)", "l(cm)", "length", "length (cm)", "长"):
            idx.setdefault("length", i)
            continue
        if h in ("w cm", "w (cm)", "w(cm)", "width", "width (cm)", "宽"):
            idx.setdefault("width", i)
            continue
        if h in ("h cm", "h (cm)", "h(cm)", "height", "height (cm)", "高"):
            idx.setdefault("height", i)
            continue
        if "dimension" in h or "measurement" in h:
            idx.setdefault("dimension", i)
            continue
        if any(
            k in h
            for k in (
                "gross kg / carton",
                "gross kg/carton",
                "gross wt",
                "gross weight",
                "毛重",
            )
        ):
            idx.setdefault("weight", i)
            continue
        if "net wt" in h or "net weight" in h or h == "净重":
            idx.setdefault("net_weight", i)
            continue
        if "total gross" in h:
            idx.setdefault("total_weight", i)
            continue
    return idx


def _to_float(v: Any) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(",", "")
    s = re.sub(r"[^\d.\-eE]", "", s)
    if not s or s in ("-", "."):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def extract_from_rows(
    rows: List[List[Any]], *, source: str
) -> Dict[str, Any]:
    hdr_i = None
    for i, row in enumerate(rows[:40]):
        if _is_header_row(row):
            hdr_i = i
            break
    if hdr_i is None:
        return {"ok": False, "materials": [], "reason": "no_header_row", "header_row": None}

    headers = [str(c or "").strip() for c in rows[hdr_i]]
    idx = _map_headers(headers)
    mats: List[Dict[str, Any]] = []
    for rno, row in enumerate(rows[hdr_i + 1 :], start=hdr_i + 2):
        cells = list(row) + [None] * max(0, len(headers) - len(row))

        def cell(key: str) -> Any:
            j = idx.get(key)
            return cells[j] if j is not None and j < len(cells) else None

        name = str(cell("name") or "").strip()
        rid_raw = str(cell("id") or "").strip()
        low_name = name.lower()
        low_id = rid_raw.lower()
        if low_id in ("total", "totals", "合计", "总计", "小计", "汇总") or low_name in (
            "total",
            "totals",
            "合计",
            "总计",
            "小计",
        ):
            break
        if re.match(r"^\d+\s+carton", low_name):
            continue
        if not name:
            continue
        # 空白模板占位行
        if not any(
            (cell(k) not in (None, "", 0, 0.0))
            for k in ("dimension", "length", "width", "height", "weight", "total_weight")
        ) and _to_float(cell("quantity")) in (None, 0):
            continue

        lwh = None
        if idx.get("length") is not None and idx.get("width") is not None and idx.get("height") is not None:
            lv, wv, hv = _to_float(cell("length")), _to_float(cell("width")), _to_float(cell("height"))
            if lv and wv and hv:
                # 列头已写 cm
                hdr_l = _norm(headers[idx["length"]])
                unit = "cm" if "cm" in hdr_l else ("mm" if "mm" in hdr_l else "")
                lwh = (_scale_to_mm(lv, unit or "cm"), _scale_to_mm(wv, unit or "cm"), _scale_to_mm(hv, unit or "cm"))
        if lwh is None:
            lwh = _parse_dim_cell(cell("dimension"))
        L, W, H = lwh if lwh else (0.0, 0.0, 0.0)

        qty = _to_float(cell("quantity"))
        qty_i = max(1, int(qty)) if qty and qty > 0 else 1
        unit_w = _to_float(cell("weight")) or 0.0
        total_w = _to_float(cell("total_weight"))
        wt_hdr = _norm(headers[idx["weight"]]) if idx.get("weight") is not None else ""
        per_piece = any(k in wt_hdr for k in ("per carton", "/ carton", "/carton", "单重"))
        # 出口箱单 Gross Wt 常是行合计；仅标明 per carton 才当单箱重
        if total_w is None and unit_w:
            if per_piece or qty_i == 1:
                total_w = unit_w * qty_i
            else:
                total_w = unit_w
                unit_w = total_w / qty_i
        if unit_w <= 0 and total_w:
            unit_w = total_w / qty_i
        if L <= 0 and W <= 0 and H <= 0 and unit_w <= 0 and not name:
            continue

        rid = str(cell("id") or f"M{len(mats)+1:03d}").strip() or f"M{len(mats)+1:03d}"
        mats.append(
            {
                "id": rid,
                "name": name,
                "quantity": qty_i,
                "weight_kg": round(float(unit_w or 0), 4),
                "total_weight_kg": round(float(total_w or 0), 4),
                "length_mm": round(float(L), 3),
                "width_mm": round(float(W), 3),
                "height_mm": round(float(H), 3),
                "category": "carton",
                "note": source,
                "meta": {
                    "source": source,
                    "header_row": hdr_i + 1,
                    "column_map": {k: headers[v] for k, v in idx.items()},
                    "dims_estimated": L <= 0 or W <= 0 or H <= 0,
                },
            }
        )
    return {
        "ok": bool(mats),
        "materials": mats,
        "reason": "" if mats else "no_data_rows",
        "header_row": hdr_i + 1,
        "headers": headers,
        "column_map": {k: headers[v] for k, v in idx.items()},
    }


def extract_piston_example(path: Path) -> Dict[str, Any]:
    """包装承认书不是表格式箱单，按样例格点抠 27 箱。"""
    name, rows = _sheet_rows(path, "Production Packaging Example")
    lwh_in = None
    n_per_pallet = 27
    full_lb = 34.0
    for row in rows:
        blob = _blob(row).lower()
        if "unit of measure" in blob and "inches" in blob:
            nums = [c for c in row if isinstance(c, (int, float)) and c > 0]
            if len(nums) >= 3:
                lwh_in = [float(nums[0]), float(nums[1]), float(nums[2])]
        if "containers per pallet" in blob:
            nums = [c for c in row if isinstance(c, (int, float)) and c > 0]
            if nums:
                n_per_pallet = int(nums[0])
        if str(row[0] or "").strip() == "Full Container":
            nums = [c for c in row if isinstance(c, (int, float)) and c > 0]
            if nums:
                full_lb = float(nums[0])
    if not lwh_in:
        return {"ok": False, "materials": [], "reason": "piston_lwh_not_found"}
    L, W, H = [_scale_to_mm(v, "in") for v in lwh_in]
    kg = full_lb * 0.453592
    mat = {
        "id": "WIDGET-CTN",
        "name": "Widget HSC carton (Piston example)",
        "quantity": n_per_pallet,
        "weight_kg": round(kg, 4),
        "total_weight_kg": round(kg * n_per_pallet, 4),
        "length_mm": round(L, 3),
        "width_mm": round(W, 3),
        "height_mm": round(H, 3),
        "category": "carton",
        "note": "piston Production Packaging Example",
        "meta": {
            "source": "piston",
            "sheet": name,
            "lwh_in": lwh_in,
            "full_container_lb": full_lb,
            "dims_estimated": False,
        },
    }
    return {"ok": True, "materials": [mat], "reason": "", "header_row": None}


def native_parse(path: Path) -> Dict[str, Any]:
    from packing_assistant.tools.table_mapper import parse_table_file

    suf = path.suffix.lower()
    if suf == ".xls":
        return {"ok": False, "materials": [], "errors": ["native parse_table_file does not load .xls"]}
    if suf == ".pdf":
        return {"ok": False, "materials": [], "errors": ["native parse_table_file does not load .pdf"]}
    try:
        return parse_table_file(path)
    except Exception as e:
        return {"ok": False, "materials": [], "errors": [f"{type(e).__name__}: {e}"]}


def _complete(m: Dict[str, Any]) -> bool:
    return (
        float(m.get("length_mm") or 0) > 0
        and float(m.get("width_mm") or 0) > 0
        and float(m.get("height_mm") or 0) > 0
    )


def _cbm(mats: Sequence[Dict[str, Any]]) -> float:
    tot = 0.0
    for m in mats:
        q = float(m.get("quantity") or 1)
        tot += (
            float(m.get("length_mm") or 0)
            * float(m.get("width_mm") or 0)
            * float(m.get("height_mm") or 0)
            * q
            / 1e9
        )
    return round(tot, 4)


def explode_packages(mats: Sequence[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """箱单一行多箱 → 一件一箱。crate_passthrough 本身是一行一箱，不炸开会少装。"""
    out: List[Dict[str, Any]] = []
    for m in mats:
        q = max(int(m.get("quantity") or 1), 1)
        unit = float(m.get("weight_kg") or 0)
        if unit <= 0 and q:
            unit = float(m.get("total_weight_kg") or 0) / q
        for i in range(q):
            item = dict(m)
            item["id"] = f"{m.get('id')}-{i+1:03d}" if q > 1 else str(m.get("id"))
            item["quantity"] = 1
            item["weight_kg"] = round(unit, 4)
            item["total_weight_kg"] = round(unit, 4)
            out.append(item)
    return out


def run_pack(label: str, mats: List[Dict[str, Any]]) -> Dict[str, Any]:
    from packing_assistant.harness import apply_user_confirmation, run_team_a, run_team_b
    from packing_assistant.pack_profile import apply_pack_profile

    opts = apply_pack_profile(
        {"profile_id": "generic_table", "crate_passthrough": True, "pack_profile": "demo"}
    )
    t0 = time.perf_counter()
    a = run_team_a(label, materials=mats, session_id=f"webpl-{label[:20]}", packing_options=opts)
    dt_a = time.perf_counter() - t0
    t1 = time.perf_counter()
    b = run_team_b(apply_user_confirmation(a, action="confirm", container_type="40HQ"))
    dt_b = time.perf_counter() - t1
    plan = b.get("container_plan") or {}
    return {
        "team_a_s": round(dt_a, 3),
        "team_b_s": round(dt_b, 3),
        "phase": b.get("phase"),
        "boxes": len(b.get("boxes") or a.get("boxes") or []),
        "structure_skipped": bool(a.get("structure_skipped")),
        "n0": plan.get("n0"),
        "used": plan.get("containers_used"),
        "can_fit": plan.get("can_fit"),
        "ship_ok": b.get("ship_ok"),
        "mid50": plan.get("worst_mid50"),
        "vol": plan.get("space_utilization"),
        "wt": plan.get("weight_utilization"),
        "pack_profile": (b.get("packing_options") or {}).get("pack_profile") or plan.get("pack_profile"),
    }


def judge(src: Dict[str, Any], adapted: Dict[str, Any], pack: Optional[Dict[str, Any]]) -> str:
    truth = src.get("doc_truth") or {}
    kind = src.get("kind")
    mats = adapted.get("materials") or []
    if kind == "blank_template":
        return "PASS-BLANK" if not mats else "UNEXPECTED-ROWS"
    if kind == "not_packing_list":
        return "SKIP-NOT-PL"
    if kind == "filled_packing_list_no_dims":
        pkg = truth.get("packages")
        got = sum(int(m.get("quantity") or 1) for m in mats)
        if mats and pkg and got == pkg:
            return "PASS-PARSE-NO-DIMS"
        if mats:
            return "PARTIAL-PARSE-NO-DIMS"
        return "FAIL-PARSE"
    if not mats:
        return "FAIL-PARSE"
    if truth.get("packages") and sum(int(m.get("quantity") or 1) for m in mats) != truth["packages"]:
        return "MISMATCH-QTY"
    if not all(_complete(m) for m in mats) and kind != "filled_packing_list_no_dims":
        return "PARSE-MISSING-LWH"
    if not pack:
        return "NO-PACK"
    exp = truth.get("expected_containers")
    if exp is not None:
        if pack.get("can_fit") and pack.get("used") == exp:
            return "PASS"
        if pack.get("can_fit") and pack.get("used") and pack.get("used") <= exp:
            return "PASS"
        return "MISMATCH-CONTAINERS"
    return "PASS" if pack.get("can_fit") else "FAIL-PACK"


def main() -> int:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    results: List[Dict[str, Any]] = []
    print("=" * 76)
    print("网上公开箱单实测（下载缓存 + 原生解析 + 箱单适配 + demo 装箱）")
    print("=" * 76)

    for src in SOURCES:
        print("-" * 76)
        print(f"[{src['id']}] {src['url']}")
        rec: Dict[str, Any] = {"id": src["id"], "url": src["url"], "kind": src["kind"]}
        try:
            path = _download(src)
            rec["bytes"] = path.stat().st_size
            rec["path"] = str(path.relative_to(ROOT))
        except Exception as e:
            rec["status"] = "DOWNLOAD-FAIL"
            rec["error"] = str(e)
            print("  DOWNLOAD FAIL", e)
            results.append(rec)
            continue

        native = native_parse(path)
        rec["native_ok"] = bool(native.get("ok"))
        rec["native_rows"] = len(native.get("materials") or [])
        rec["native_errors"] = native.get("errors") or []
        rec["native_stats"] = native.get("stats")

        adapted: Dict[str, Any]
        if src["id"] == "piston_widget_example":
            adapted = extract_piston_example(path)
        elif path.suffix.lower() == ".pdf":
            try:
                from packing_assistant.tools.packing_list_parser import parse_packing_list_pdf

                pdf = parse_packing_list_pdf(path)
                adapted = {
                    "ok": bool(pdf.get("materials")),
                    "materials": pdf.get("materials") or [],
                    "reason": "pdf_parser",
                    "header_row": None,
                    "pdf_meta": {k: pdf.get(k) for k in ("contract", "n_lines") if k in pdf},
                }
            except Exception as e:
                adapted = {"ok": False, "materials": [], "reason": f"pdf:{e}"}
        else:
            _sheet, rows = _sheet_rows(path, src.get("sheet"))
            rec["sheet"] = _sheet
            adapted = extract_from_rows(rows, source=src["id"])

        mats = list(adapted.get("materials") or [])
        rec["adapted_ok"] = bool(adapted.get("ok"))
        rec["adapted_reason"] = adapted.get("reason")
        rec["adapted_header_row"] = adapted.get("header_row")
        rec["adapted_map"] = adapted.get("column_map")
        rec["n_lines"] = len(mats)
        rec["n_packages"] = sum(int(m.get("quantity") or 1) for m in mats)
        rec["weight_kg"] = round(sum(float(m.get("total_weight_kg") or 0) for m in mats), 3)
        rec["cbm"] = _cbm(mats)
        rec["complete_lwh"] = sum(1 for m in mats if _complete(m))
        rec["sample"] = [
            {
                "id": m.get("id"),
                "name": (m.get("name") or "")[:60],
                "qty": m.get("quantity"),
                "lwh_mm": [m.get("length_mm"), m.get("width_mm"), m.get("height_mm")],
                "kg": m.get("weight_kg"),
            }
            for m in mats[:4]
        ]

        pack = None
        packable = [m for m in mats if _complete(m)]
        exploded = explode_packages(packable) if packable else []
        rec["exploded_packages"] = len(exploded)
        rec["passthrough_would_be_boxes"] = len(packable)
        if exploded:
            try:
                pack = run_pack(src["id"], exploded)
                rec["pack"] = pack
            except Exception as e:
                rec["pack_error"] = f"{type(e).__name__}: {e}"
        rec["status"] = judge(src, adapted, pack)
        rec["doc_truth"] = src.get("doc_truth")

        print(
            f"  native ok={rec['native_ok']} rows={rec['native_rows']}  "
            f"adapt n={rec['n_lines']} pkg={rec['n_packages']} "
            f"lwh={rec['complete_lwh']} kg={rec['weight_kg']} cbm={rec['cbm']}"
        )
        if pack:
            print(
                f"  pack n0={pack.get('n0')} used={pack.get('used')} "
                f"fit={pack.get('can_fit')} ship={pack.get('ship_ok')} "
                f"mid50={pack.get('mid50')} vol={pack.get('vol')} "
                f"A={pack.get('team_a_s')}s B={pack.get('team_b_s')}s"
            )
        print(f"  truth: {src['doc_truth'].get('note')}")
        print(f"  → {rec['status']}")
        results.append(rec)

    summary = {
        "title": "Public internet packing lists vs packing-agent",
        "n": len(results),
        "by_status": {},
        "results": results,
    }
    for r in results:
        st = r.get("status") or "?"
        summary["by_status"][st] = summary["by_status"].get(st, 0) + 1
    out = OUT_DIR / "web_real_packing_lists.json"
    out.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print("=" * 76)
    print("status counts:", summary["by_status"])
    print("report:", out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
