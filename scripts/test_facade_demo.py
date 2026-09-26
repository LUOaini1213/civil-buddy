#!/usr/bin/env python3
"""scripts/demo_facade.py on the SYNTHETIC façade pack: the three flows run offline and say what they found.

  tender     CR16, 420 calendar days, 90-day validity, the 10% bond and the 12-month DLP land in their rows;
             the two bid posts write from the same session's hand-off
  packing    24 panels / 10,800 kg in the list are 24 / 10,800 kg in the crates, and the plan fits
  site docs  a daily report is written; the work-at-height briefing (high risk) writes nothing until the
             person's sentence is given, and is written once it is
  scope      nothing is written outside the job folder the demo creates; a wrong --sign or a used folder
             is refused before any turn runs
No model and no network: steps mode only.
"""
from __future__ import annotations

import contextlib
import importlib.util
import io
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
os.environ["PYTHON_DOTENV_DISABLED"] = "1"
for _key in [k for k in os.environ if k.endswith("_API_KEY")] + ["CIVIL_AGENT_MODE", "CIVIL_SANDBOX", "CIVIL_APPROVAL"]:
    os.environ.pop(_key, None)

spec = importlib.util.spec_from_file_location("demo_facade", ROOT / "scripts" / "demo_facade.py")
demo = importlib.util.module_from_spec(spec)
spec.loader.exec_module(demo)

from packing_assistant.civil import CONFIRM  # noqa: E402
from packing_assistant.runtime import workspace  # noqa: E402

_WRITE_FLAGS = os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_APPEND | os.O_TRUNC
_WATCH = {"on": False, "writes": []}


def _audit(event, args):
    if not _WATCH["on"]:
        return
    path = None
    if event == "open" and args and isinstance(args[0], (str, bytes, os.PathLike)):
        mode, flags = (args[1] if len(args) > 1 else None), (args[2] if len(args) > 2 else 0)
        if (isinstance(mode, str) and any(c in mode for c in "wax+")) or (isinstance(flags, int) and flags & _WRITE_FLAGS):
            path = args[0]
    elif event in {"os.mkdir", "os.rename", "os.replace", "os.remove", "os.rmdir", "shutil.copyfile", "shutil.rmtree"} and args:
        path = args[1] if event in {"os.rename", "os.replace", "shutil.copyfile"} and len(args) > 1 else args[0]
    elif event == "sqlite3.connect" and args:
        path = args[0]
    if path is not None and not isinstance(path, int):
        _WATCH["writes"].append(os.fsdecode(path))


sys.addaudithook(_audit)


class FacadeDemo(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix="facade-demo-test-")
        cls.base = Path(cls.tmp.name).resolve()
        cls.job = cls.base / "job"
        cls.cwd = Path.cwd()
        cls.out = io.StringIO()
        home = patch.object(Path, "home", return_value=cls.base / "no-home")     # the user's ~/.civil-buddy stays out
        home.start()
        cls.addClassCleanup(home.stop)
        mpl = os.environ.get("MPLCONFIGDIR")
        _WATCH.update(on=True, writes=[])
        try:
            with contextlib.redirect_stdout(cls.out):
                cls.result = demo.run_demo(cls.job)
                cls.signed = demo.Demo(cls.job, sign=CONFIRM)     # the person types the sentence
                cls.signed.briefing()
        finally:
            _WATCH["on"] = False
            workspace.deactivate()
            os.chdir(cls.cwd)
            tempfile.tempdir = None
            if mpl is None:
                os.environ.pop("MPLCONFIGDIR", None)
            else:
                os.environ["MPLCONFIGDIR"] = mpl
        sys.stdout.write(cls.out.getvalue()[-4000:])

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_no_flow_errored(self):
        self.assertEqual(self.result["errors"], [])
        self.assertEqual(self.signed.errors, [])

    def test_tender_rows_and_bid_posts(self):
        rows = self.result["tender"]["rows"]
        self.assertIn("CR16", rows["注册资格/工作类别"])
        self.assertIn("420 calendar days", rows["工期"])
        self.assertIn("90 days", rows["投标有效期"])
        self.assertIn("10%", rows["履约担保"])
        self.assertIn("12 months", rows["缺陷责任期/质保期"])
        self.assertEqual(len(self.result["tender"]["scoring"]), 4)
        self.assertTrue(self.result["tender"]["submit_blocked"])
        self.assertTrue(self.result["bid_tech"]["written"])
        self.assertTrue(self.result["bid_compliance"]["written"])
        self.assertTrue((self.job / ".civil-buddy" / "out" / "civil-cli" / "bid-parse" / "tender.parse.md").is_file())

    def test_panels_are_conserved_and_fit(self):
        packing = self.result["packing"]
        cons = packing["conservation"]
        self.assertTrue(cons["ok"], cons)
        self.assertEqual((cons["pieces_in"], cons["pieces_out"]), (24, 24))
        self.assertEqual((cons["kg_in"], cons["kg_out"]), (10800, 10800))
        self.assertIs(packing["plan"]["can_fit"], True)
        self.assertGreaterEqual(packing["plan"]["containers_used"], 1)
        self.assertIn("pieces 24 -> 24 · kg 10800.0 -> 10800.0 · ok", self.out.getvalue())

    def test_daily_report_is_written(self):
        daily = self.result["daily"]
        self.assertTrue(daily["written"])
        self.assertEqual(daily["rows"]["日期"], "2026-09-24")
        self.assertEqual(daily["rows"]["部位"], "东立面五层")

    def test_briefing_waits_for_the_person(self):
        self.assertTrue(self.result["briefing"]["refused_without_sentence"])
        self.assertFalse(self.result["briefing"]["written"])
        self.assertIn("This demo does not supply the sentence", self.out.getvalue())
        self.assertIn(CONFIRM, self.out.getvalue())
        signed = self.signed.result["briefing"]
        self.assertTrue(signed["refused_without_sentence"])      # the signed run still asks first
        self.assertTrue(signed["written"])
        self.assertTrue((self.job / signed["file"]).is_file())

    def test_nothing_is_written_outside_the_job_folder(self):
        outside = sorted({p for p in _WATCH["writes"] if not Path(p).resolve().is_relative_to(self.job)})
        self.assertEqual(outside, [], "\n".join(outside))
        self.assertGreater(len(_WATCH["writes"]), 10)          # the hook did see the drafts being written

    def test_fixtures_say_they_are_synthetic(self):
        import openpyxl

        for name in ("facade_itt_doc.md", "daily_report_input.txt", "wah_briefing_input.txt", "README.md"):
            head = (demo.FIXTURES / name).read_text(encoding="utf-8")[:400]
            self.assertRegex(head, r"SYNTHETIC|合成示例", name)
        for name in ("facade_panels.xlsx", "facade_panels_zh.xlsx"):
            wb = openpyxl.load_workbook(demo.FIXTURES / name, read_only=True)
            try:
                self.assertEqual(wb.sheetnames, ["materials", "README"])
                self.assertIn("SYNTHETIC", str(next(wb["README"].iter_rows(values_only=True))[0]))
                names = [row[1] for row in wb["materials"].iter_rows(min_row=2, values_only=True)]
                self.assertTrue(names and all("SYNTHETIC" in n or "合成示例" in n for n in names), names)
            finally:
                wb.close()

    def test_bad_arguments_are_refused_before_any_turn(self):
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(demo.main(["--sign", "yes"]), 2)
            self.assertEqual(demo.main(["--job", str(self.job)]), 2)         # used folder: never written into
        self.assertIsNone(workspace.active())


if __name__ == "__main__":
    unittest.main(verbosity=1)
