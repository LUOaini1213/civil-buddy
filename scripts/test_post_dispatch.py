#!/usr/bin/env python3
"""Post writers must share persistence, truthful failures and chat isolation."""
from __future__ import annotations

import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from packing_assistant import expert_turn as posts
from packing_assistant.runtime.scheduler import Scheduler
from packing_assistant.runtime.memory import assemble_context as assemble_session_context
from packing_assistant.runtime.tool_engine import ToolEngine, default_engine


class PostDispatchTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="civil-post-dispatch-")
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.scheduler = Scheduler()
        for context in (
            patch.object(posts, "_OUT", self.root),
            patch.dict(os.environ, {"CIVIL_SANDBOX_ROOTS": str(self.root), "CIVIL_JOB_ROOT": "", "PYTHON_DOTENV_DISABLED": "1"}),
            patch("packing_assistant.runtime.scheduler.get_scheduler", return_value=self.scheduler),
            patch("packing_assistant.runtime.memory.assemble_context", return_value={}),
        ):
            context.start()
            self.addCleanup(context.stop)

    def _run(self, text="写一份日报模板", post="pm-daily", sid="post-test"):
        return posts.run_expert_turn(text, post, session_id=sid, confirm_ok=True, force_intent="run")

    def test_all_migrated_renderers_retain_exact_document_and_tool_contract(self):
        for post, (builder, tool, reply) in posts._SIMPLE_DRAFTS.items():
            with self.subTest(post=post):
                text = "写一份内部讨论模板；辖区：CN"
                result = self._run(text, post, post)
                self.assertTrue(result["ok"], result)
                self.assertEqual(result["state"], "done")
                markdown = next(f for f in result["files"] if f["tool"] == tool)
                self.assertEqual(Path(markdown["path"]).read_text(encoding="utf-8"), builder(text))
                self.assertTrue(result["reply"].startswith(reply))
                self.assertTrue(result["submit_blocked"])

    def test_chat_cannot_execute_written_checklists(self):
        engine = default_engine()
        for post, tool in (("admin-office", "admin-office__list"), ("lab-sample", "lab-sample__list"), ("env", "env__list")):
            with self.subTest(tool=tool):
                self.assertTrue(engine.tools[tool].writes)
                result = engine.execute(tool, {"text": "写一份清单", "session_id": "no-write"}, expert_id=post, intent="chat")
                self.assertFalse(result["ok"])
                self.assertEqual(result["error_code"], "permission_denied")
        self.assertFalse(engine.tools["pack-ship__list"].writes)
        self.assertFalse(list(self.root.rglob("*.md")))

    def test_office_failure_retains_other_real_formats_and_reports_failure(self):
        for failed_format, retained_suffix in (("docx", ".xlsx"), ("xlsx", ".docx")):
            with self.subTest(failed_format=failed_format), patch(
                "packing_assistant.office_job.export_md_to_" + failed_format,
                side_effect=OSError("fixture export failure"),
            ):
                result = self._run(sid="office-failed-" + failed_format)
                self.assertFalse(result["ok"], result)
                self.assertEqual(result["state"], "failed")
                self.assertEqual(result["error_code"], "office_export_failed")
                self.assertTrue(result["wrote"])
                self.assertTrue(any(Path(f["path"]).suffix == retained_suffix for f in result["files"]))
                self.assertTrue(all(Path(f["path"]).is_file() for f in result["files"]))

    def test_non_boolean_confirmation_never_authorizes_high_risk_write(self):
        for value in ("false", "true", 1, [True], {"confirmed": True}):
            with self.subTest(value=value):
                result = posts.run_named_exclusive("safety-brief__talk", {
                    "text": "写一份安全交底", "session_id": "typed-confirm", "confirm_ok": value,
                    "p0_confirmed": value,
                })
                self.assertTrue(result["hitl_pending"])
                self.assertFalse(result["wrote"])
                self.assertFalse(result["files"])

    def test_read_only_blocks_engine_and_direct_expert_writes(self):
        with patch.dict(os.environ, {"CIVIL_SANDBOX": "read-only"}):
            engine = default_engine()
            result = engine.execute("admin-office__list", {"text": "写一份清单", "session_id": "read-only"}, expert_id="admin-office")
            direct = self._run(post="admin-office", sid="read-only")
        for item in (result, direct):
            self.assertFalse(item["ok"])
            self.assertEqual(item["error_code"], "permission_denied")
            self.assertFalse(item.get("files"))
        self.assertFalse(list(self.root.rglob("*.md")))

    def test_string_confirmation_in_memory_is_not_trusted(self):
        with patch("packing_assistant.runtime.memory.load_summary", return_value={"p0_confirmed": "false"}), patch("packing_assistant.runtime.memory.save_summary"):
            context = assemble_session_context("bad-slot", p0_confirmed="true")
        self.assertFalse(context["p0_confirmed"])

    def test_all_documents_are_scanned_before_first_write(self):
        result = posts._save_drafts(self.root, [("one", "# 内部草稿"), ("two", "可以开工")], "saved")
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "forbidden_content")
        self.assertFalse(list(self.root.iterdir()))

    def test_partial_write_failure_preserves_only_actual_files(self):
        real_write = posts.guarded_write_text

        def save(path, text):
            if path.name == "second.md":
                raise OSError("test disk full")
            return real_write(path, text)

        with patch.object(posts, "guarded_write_text", side_effect=save):
            result = posts._save_drafts(self.root, [("first", "已写入"), ("second", "未写入")], "saved")
        self.assertFalse(result["ok"])
        self.assertTrue(result["wrote"])
        self.assertEqual(result["error_code"], "write_failed")
        self.assertEqual([f["name"] for f in result["files"]], ["first.md"])
        self.assertEqual((self.root / "first.md").read_text(encoding="utf-8"), "已写入")
        self.assertFalse((self.root / "second.md").exists())

    def test_structured_tool_failure_is_not_wrapped_as_success(self):
        engine = ToolEngine()
        engine.register("broken", lambda _: {"ok": False, "error_code": "write_failed", "files": [], "reply": "failed"})
        result = engine.execute("broken")
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "write_failed")
        self.assertEqual(engine.audit_log[-1].error_code, "write_failed")

    def test_excel_decodes_rendered_cells_once_without_executing_formulas(self):
        from packing_assistant.office_job import export_md_to_xlsx
        import openpyxl

        path = self.root / "entities.md"
        path.write_text("| 条件 | 值 |\n| --- | --- |\n| 小于 | x &lt; 10 |\n| 管道 | a&#124;b |\n| 原始实体 | &amp;lt; |\n| 文本 | &#61;1+1 |\n", encoding="utf-8")
        exported = export_md_to_xlsx(path)
        workbook = openpyxl.load_workbook(exported[0], data_only=False)
        try:
            sheet = workbook.worksheets[0]
            self.assertEqual(sheet.cell(2, 2).value, "x < 10")
            self.assertEqual(sheet.cell(3, 2).value, "a|b")
            self.assertEqual(sheet.cell(4, 2).value, "&lt;")
            self.assertEqual(sheet.cell(5, 2).value, "=1+1")
            self.assertEqual(sheet.cell(5, 2).data_type, "s")
        finally:
            workbook.close()

    def test_forbidden_content_marks_expert_run_failed(self):
        with patch.dict(posts._SIMPLE_DRAFTS, {"pm-daily": (lambda _: "可以开工", "pm-daily__log", "saved")}):
            result = self._run()
        self.assertFalse(result["ok"])
        self.assertEqual(result["state"], "failed")
        self.assertEqual(result["error_code"], "forbidden_content")
        self.assertFalse(result["files"])

    def test_exception_releases_session_for_retry(self):
        with patch.object(posts, "_run_exclusive", side_effect=RuntimeError("test builder exception")):
            failed = self._run()
        self.assertFalse(failed["ok"])
        self.assertEqual(failed["state"], "failed")
        self.assertNotIn("test builder exception", failed["reply"])
        retried = self._run()
        self.assertTrue(retried["ok"], retried)
        self.assertEqual(retried["state"], "done")

    def test_busy_run_does_not_write_or_release_active_owner(self):
        active = self.scheduler.create_run("post-test")
        with patch.object(posts, "_run_exclusive") as writer:
            result = self._run()
        self.assertFalse(result["ok"])
        self.assertEqual(result["error_code"], "session_busy")
        writer.assert_not_called()
        self.assertEqual(self.scheduler.create_run("post-test").error_code, "session_busy")
        self.assertEqual(active.state, "pending")


if __name__ == "__main__":
    unittest.main()
