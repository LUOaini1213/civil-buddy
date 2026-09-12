#!/usr/bin/env python3
"""Administration drafts: structured facts, approval boundaries and real Office output."""

from __future__ import annotations

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from packing_assistant.expert_roster import exclusive_tools
from packing_assistant.office_job import tables_from_md
from packing_assistant.post_drafts.admin import build_draft


def draft(post: str, text: str) -> str:
    value = build_draft(post, exclusive_tools(post)[0], text)
    assert isinstance(value, str)
    return value


def rows_for(markdown: str, title: str) -> list[list[str]]:
    return next(rows for name, rows in tables_from_md(markdown) if name.startswith(title))


class AdminDraftTests(unittest.TestCase):
    def test_roster_tools_and_mismatches(self) -> None:
        for post in ("admin-doc", "admin-office"):
            for tool in exclusive_tools(post):
                with self.subTest(post=post, tool=tool):
                    self.assertIsInstance(build_draft(post, tool, "写一份草稿"), str)
        self.assertIsNone(build_draft("admin-doc", "admin-office__list", "写请示"))
        self.assertIsNone(build_draft("admin-office", "unknown", "写会务清单"))
        self.assertIsNone(build_draft("hr-train", "hr-train__plan", "写培训计划"))

    def test_request_populates_facts_and_preserves_document_number(self) -> None:
        md = draft("admin-doc", "写一份请示\n文种：请示；标题：关于补配会议设备的请示；"
                   "发文机关：项目综合部；文号：综〔2026〕8号；主送：公司办公室；"
                   "请示事项：补配投影设备；背景：现有投影无法开机；依据：用户提供的设备台账；"
                   "请求事项：请批准设备检修；附件：故障记录、报价材料；成文日期：2026-09-12")
        self.assertIn(["事实与背景", "现有投影无法开机"], rows_for(md, "请示正文"))
        self.assertIn(["请求事项", "请批准设备检修"], rows_for(md, "请示正文"))
        self.assertIn(["主送", "公司办公室"], rows_for(md, "文件控制"))
        self.assertIn(["发文字号（仅录入原值，不生成序号）", "综〔2026〕8号"], rows_for(md, "文件控制"))
        self.assertIn(["批示意见", ""], rows_for(md, "请示正文"))
        self.assertIn("妥否，请批示", md)
        self.assertNotIn("## 用户原文", md)

    def test_minutes_do_not_promote_discussion_or_proposals(self) -> None:
        md = draft("admin-doc", "写会议纪要；会议名称：材料协调会；讨论：是否调整供货计划；"
                   "决议：建议下周再讨论；待定事项：供货日期；责任人：张工；期限：下周")
        decisions = rows_for(md, "已议定事项")
        self.assertTrue(decisions[1][0].startswith("UNSPECIFIED"))
        self.assertTrue(decisions[1][1].startswith("UNSPECIFIED"))
        self.assertIn("建议下周再讨论", rows_for(md, "待定事项")[1][0])
        self.assertEqual("", rows_for(md, "待定事项")[1][-1])
        self.assertIn(["讨论摘要", "是否调整供货计划"], rows_for(md, "会议记录要素"))

    def test_minutes_transfer_only_explicit_decision_with_owner_and_deadline(self) -> None:
        md = draft("admin-doc", "文种：会议纪要；会议名称：办公室周会；时间：2026-09-12 09:00；"
                   "地点：302室；主持人：陈主任；与会人员：综合部、设备部；记录人：小刘；"
                   "决议：补交设备故障记录；责任人：张工；截止时间：2026-09-15")
        self.assertEqual(["补交设备故障记录", "张工", "2026-09-15", ""], rows_for(md, "已议定事项")[1])
        self.assertIn(["与会人员", "综合部、设备部"], rows_for(md, "会议记录要素"))
        self.assertIn(["记录人", "小刘"], rows_for(md, "会议记录要素"))

    def test_labeled_proposal_is_still_pending(self) -> None:
        for proposal in ("拟调整会议安排", "预算待批准", "讨论更换会议场地"):
            with self.subTest(proposal=proposal):
                md = draft("admin-doc", f"写会议纪要；决议：{proposal}")
                self.assertTrue(rows_for(md, "已议定事项")[1][0].startswith("UNSPECIFIED"))
                self.assertEqual(proposal, rows_for(md, "待定事项")[1][0])

    def test_seal_form_does_not_sign_or_execute(self) -> None:
        md = draft("admin-doc", "写用印申请；文件名称：场地租赁说明；份数：2份；印章种类：公司公章；"
                   "审批人：王主任；申请人：陈工；用印事由：提交场地使用材料；"
                   "送达对象：场地方；是否骑缝：待核；附件：无")
        rows = rows_for(md, "用印申请与核验")
        for row in (["文件名称", "场地租赁说明"], ["份数", "2份"], ["印章种类", "公司公章"],
                    ["拟审批人（用户提供，权限待核）", "王主任"], ["审批人签认", ""],
                    ["实际用印日期及登记编号", ""]):
            self.assertIn(row, rows)
        self.assertIn("禁止代用印、预盖空白文件", md)
        self.assertIn(["附件目录（按用户顺序）", "无"], rows_for(md, "附件与归档交接"))

    def test_missing_or_ambiguous_document_kind_does_not_invent_a_request(self) -> None:
        for text in ("写一份公文", "写请示报告", "文种：通知；主题：设备盘点"):
            with self.subTest(text=text):
                md = draft("admin-doc", text)
                self.assertIn("文种待确认", md)
                self.assertNotIn("## 请示正文", md)
                self.assertIn("UNSPECIFIED", rows_for(md, "附件与归档交接")[1][1])

    def test_office_four_tables_populate_inputs_and_keep_decisions_blank(self) -> None:
        md = draft("admin-office", "写会务清单；会议名称：周协调会；目的：梳理资料缺项；"
                   "时间：周五 14:00；场地：项目部二楼；议程：资料盘点、责任分工；"
                   "主持人：李工；与会人员：工程部、综合部；资料目录：上周纪要、图纸目录；"
                   "设备：投影仪；决议：同意所有安排")
        self.assertEqual("项目部二楼", rows_for(md, "场地")[1][1])
        self.assertEqual("资料盘点、责任分工", rows_for(md, "议程")[1][1])
        self.assertEqual("工程部、综合部", rows_for(md, "与会")[1][1])
        self.assertEqual("上周纪要、图纸目录", rows_for(md, "资料目录")[1][0])
        for title in ("场地", "议程", "与会", "资料目录"):
            rows = rows_for(md, title)
            decision = rows[0].index("决定")
            self.assertTrue(all(row[decision] == "" for row in rows[1:]), title)
        self.assertNotIn("同意所有安排", md)

    def test_common_narrative_meeting_facts_are_extracted(self) -> None:
        md = draft("admin-office", "请写会务清单，计划于9月15日14:00在项目部二楼召开周例会，"
                   "由李工主持，由小张记录，参会人员为工程部和综合部。")
        self.assertIn(["时间", "9月15日14:00"], rows_for(md, "任务信息"))
        self.assertIn(["会议/活动名称", "周例会"], rows_for(md, "任务信息"))
        self.assertEqual("项目部二楼", rows_for(md, "场地")[1][1])
        self.assertEqual("李工", rows_for(md, "议程")[1][2])
        self.assertEqual("小张", rows_for(md, "与会")[3][1])

    def test_requested_reception_travel_and_supplies_have_specific_fields(self) -> None:
        md = draft("admin-office", "编制后勤清单；接待对象：合作单位代表；人数：3人；住宿：用户待确认；"
                   "出差人员：李工；路线：上海至苏州；交通方式：铁路；"
                   "申请部门：综合部；办公物资：文件盒；规格：A4；数量：10个；库存：2个；领用人：小张；用途：资料归档")
        self.assertIn(["人数", "3人"], rows_for(md, "接待安排"))
        self.assertIn(["路线", "上海至苏州"], rows_for(md, "差旅及用车"))
        self.assertEqual(["综合部", "文件盒", "A4", "10个", "2个", "小张", "资料归档", ""], rows_for(md, "办公物资申领")[1])
        self.assertIn(["费用标准及审批", ""], rows_for(md, "接待安排"))

    def test_jurisdiction_is_shared_and_missing_values_stay_unknown(self) -> None:
        for post in ("admin-doc", "admin-office"):
            for prefix, expected in (("", "UNSPECIFIED"), ("CN ", "CN"), ("新加坡 ", "SG"), ("EU CN ", "DUAL")):
                with self.subTest(post=post, prefix=prefix):
                    md = draft(post, prefix + "写一份请示" if post == "admin-doc" else prefix + "写会务清单")
                    self.assertIn(f"- 辖区：{expected}\n", md)
                    self.assertIn("草稿声明", md)
                    self.assertIn("不代用印", md)
        self.assertIn(["签发人签认", ""], rows_for(draft("admin-doc", "写请示"), "文件控制"))

    def test_field_boundaries_and_markdown_table_cells_are_safe(self) -> None:
        md = draft("admin-office", "场地：A|B<script>；议程：资料盘点, 确认记录模板\n"
                   "额外说明不属于议程字段\n主持人：李工")
        self.assertEqual("A|B<script>", rows_for(md, "场地")[1][1])
        self.assertEqual(4, len(rows_for(md, "场地")[1]))
        self.assertEqual("资料盘点, 确认记录模板", rows_for(md, "议程")[1][1])
        self.assertNotIn("<script>", md)

    def test_key_value_markdown_table_populates_office(self) -> None:
        md = draft("admin-office", "| 栏位 | 值 |\n| --- | --- |\n| 场地 | 二楼会议室 |\n| 议程 | 资料盘点 |\n| 附件 | 图纸目录 |")
        self.assertEqual("二楼会议室", rows_for(md, "场地")[1][1])
        self.assertEqual("资料盘点", rows_for(md, "议程")[1][1])
        self.assertEqual("图纸目录", rows_for(md, "资料目录")[1][0])

    def test_column_table_keeps_meeting_rows_and_missing_values_separate(self) -> None:
        md = draft("admin-office", "| 时间 | 场地 | 议程 | 主持人 | 资料目录 |\n| --- | --- | --- | --- | --- |\n"
                   "| 上午 | 一楼 | 盘点资料 | 李工 | 图纸目录 |\n| 下午 | 二楼 | 核对设备 | | 设备台账 |")
        agenda = rows_for(md, "议程")
        self.assertEqual(["上午", "盘点资料", "李工", ""], agenda[1])
        self.assertEqual("下午", agenda[2][0])
        self.assertTrue(agenda[2][2].startswith("UNSPECIFIED"))
        self.assertEqual(["一楼", "二楼"], [row[1] for row in rows_for(md, "场地")[1:3]])
        self.assertEqual(["图纸目录", "设备台账"], [row[0] for row in rows_for(md, "资料目录")[1:3]])

    def test_separate_decisions_never_share_another_records_deadline(self) -> None:
        for text in (
            "写会议纪要\n决议：整理目录；责任人：甲；期限：周五\n责任人：乙；决议：核对台账",
            "写会议纪要\n决议：整理目录\n责任人：甲\n期限：周五\n决议：核对台账\n责任人：乙",
        ):
            with self.subTest(text=text):
                rows = rows_for(draft("admin-doc", text), "已议定事项")
                self.assertEqual(["整理目录", "甲", "周五", ""], rows[1])
                self.assertEqual("核对台账", rows[2][0])
                self.assertEqual("乙", rows[2][1])
                self.assertTrue(rows[2][2].startswith("UNSPECIFIED"))

    def test_decision_column_table_keeps_each_owner_and_deadline(self) -> None:
        md = draft("admin-doc", "写会议纪要\n| 决议 | 责任人 | 期限 |\n| --- | --- | --- |\n"
                   "| 整理目录 | 甲 | 周五 |\n| 核对台账 | 乙 | |\n| 拟调整预算 | 丙 | 周一 |")
        rows = rows_for(md, "已议定事项")
        self.assertEqual(3, len(rows))
        self.assertEqual(["整理目录", "甲", "周五", ""], rows[1])
        self.assertEqual("乙", rows[2][1])
        self.assertTrue(rows[2][2].startswith("UNSPECIFIED"))
        self.assertIn("拟调整预算", rows_for(md, "待定事项")[1][0])

    def test_seal_table_preserves_file_to_copies_association(self) -> None:
        md = draft("admin-doc", "写用印申请\n| 文件名称 | 份数 | 印章种类 |\n| --- | --- | --- |\n"
                   "| 交接单甲 | 2份 | 资料章 |\n| 交接单乙 | | |")
        rows = rows_for(md, "逐份用印文件核验")
        self.assertEqual(["交接单甲", "2份", "资料章"], rows[1][:3])
        self.assertEqual("交接单乙", rows[2][0])
        self.assertTrue(rows[2][1].startswith("UNSPECIFIED"))
        self.assertTrue(rows[2][2].startswith("UNSPECIFIED"))
        self.assertIn(["份数", "见逐份用印文件核验（不合并份数或权限）"], rows_for(md, "用印申请与核验"))


class AdminRuntimeTests(unittest.TestCase):
    def test_real_runtime_markdown_and_excel_contain_structured_facts(self) -> None:
        from packing_assistant import expert_turn
        from packing_assistant.runtime import memory
        import openpyxl

        with tempfile.TemporaryDirectory(prefix="civil-admin-") as tmp:
            base = Path(tmp)
            with patch.dict(os.environ, {"CIVIL_JOB_ROOT": str(base / "job"), "CIVIL_SANDBOX_ROOTS": str(base), "PYTHON_DOTENV_DISABLED": "1"}), \
                 patch.object(expert_turn, "_OUT", base / "sessions"), patch.object(memory, "_OUT", base / "sessions"):
                cases = (
                    ("admin-doc", "写请示；请示事项：会议设备检修；背景：设备无法开机；请求事项：批准检修", "请示正文", "设备无法开机"),
                    ("admin-doc", "写会议纪要；会议名称：测试周会；决议：补交资料目录；责任人：李工；期限：周五", "已议定事项", "补交资料目录"),
                    ("admin-doc", "写用印申请；文件名称：测试资料交接单；份数：3份；印章种类：资料章", "用印申请与核验", "测试资料交接单"),
                    ("admin-office", "写会务清单；场地：测试会议室；议程：盘点资料；资料目录：图纸清单", "场地", "测试会议室"),
                    ("admin-office", "写会务清单；场地：面积 < 20 且分区 A|B", "场地", "面积 < 20 且分区 A|B"),
                )
                for index, (post, text, table, expected) in enumerate(cases):
                    with self.subTest(post=post):
                        result = expert_turn.run_expert_turn(text, post, force_intent="run", session_id=f"admin-draft-{index}")
                        self.assertTrue(result["ok"], result)
                        self.assertTrue(result["wrote"], result)
                        self.assertEqual("done", result["state"])
                        files = [Path(item["path"]) for item in result["files"]]
                        md = next(path for path in files if path.suffix == ".md")
                        self.assertTrue(md.is_relative_to(base))
                        self.assertIn(expected, str(rows_for(md.read_text(encoding="utf-8"), table)))
                        book = next(path for path in files if path.suffix == ".xlsx")
                        self.assertTrue(book.is_relative_to(base))
                        workbook = openpyxl.load_workbook(book)
                        try:
                            sheet = next(workbook[name] for name in workbook.sheetnames if name.startswith(table))
                            self.assertTrue(any(expected in row for row in sheet.iter_rows(values_only=True)))
                            if post == "admin-office":
                                self.assertTrue(all(row[3] in (None, "") for row in sheet.iter_rows(min_row=2, values_only=True)))
                        finally:
                            workbook.close()

    def test_question_does_not_create_draft(self) -> None:
        from packing_assistant.expert_turn import run_expert_turn

        result = run_expert_turn("用印清单需要哪些信息？", "admin-doc", force_intent="chat", session_id="admin-question-test")
        self.assertTrue(result["ok"], result)
        self.assertFalse(result["wrote"])
        self.assertEqual([], result["files"])


if __name__ == "__main__":
    unittest.main(verbosity=2)
